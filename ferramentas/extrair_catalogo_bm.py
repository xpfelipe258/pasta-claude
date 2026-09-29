"""Extrai o catálogo de atividades de medição (BM) das planilhas de medição de cada empresa.

Para cada empresa lê a aba de memória de cálculo (MC: itens, atividades, pesos, quantidades) e os
valores contratuais de cada item nas abas QPC e RÓTULA. Gera programa_obra198/catalogo_bm.json.

Uso: python extrair_catalogo_bm.py BM_EJ.xlsx BM_CMM.xlsx BM_GLOBO.xlsx
"""
import json
import os
import re
import sys

import openpyxl

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SAIDA = os.path.join(RAIZ, "programa_obra198", "catalogo_bm.json")

# empresa -> (aba MC, aba QPC, aba RÓTULA, coluna do valor nas abas de BM, marca da empresa na col. B do MC)
CONFIG = {
    "EJ": ("BM2-MC MONTAGEM", "BM - EJ MONTAGENS - QPC", "BM1 - EJ MONTAGENS - ROTULA", 6, "E.J"),
    "CMM": ("BM2 - MC MONTAGEM", "BM - CMM - QPC", "BM - CMM - ROTULA", 6, "CMM"),
    "GLOBO AÇOS": ("MC MONTAGEM ", "BM GLOBO AÇOS - QPC", "BM GLOBO AÇOS - RÓTULA", 7, "GLOBO"),
}

# Atividades que já têm coluna no lançamento de produção (mesma unidade da quantidade do MC).
CONTROLE = {
    ("EJ", "3.1.1", "vigas senoidais"): "VIGAS",
    ("EJ", "3.1.1", "pré-montagem de joist"): "PREMONTAGEM",
    ("EJ", "3.1.1", "instalação de joist"): "IÇAMENTO JOIST",
    ("CMM", "3.1.1", "vigas senoidais"): "VIGAS",
    ("CMM", "3.1.1", "pré-montagem de joist"): "PREMONTAGEM",
    ("CMM", "3.1.1", "instalação de joist"): "IÇAMENTO JOIST",
}


def codigos(txt):
    return frozenset(re.findall(r"\d+(?:\.\d+)+", str(txt or "")))


def unidade(u):
    t = str(u or "").strip().lower()
    return {"und": "un", "un": "un", "un ": "un", "vãos": "vão", "vão": "vão", "rua": "rua", "ruas": "rua"}.get(t, t)


def valores_por_item(ws, col):
    saida = {}
    for r in range(1, ws.max_row + 1):
        cod = codigos(ws.cell(r, 1).value)
        v = ws.cell(r, col).value
        if cod and isinstance(v, (int, float)):
            saida[cod] = float(v)
    return saida


def achar(dic, cod):
    if cod in dic:
        return dic[cod]
    return next((v for k, v in dic.items() if k & cod), None)


def extrair(empresa, caminho):
    aba_mc, aba_qpc, aba_rot, col_valor, marca = CONFIG[empresa]
    wb = openpyxl.load_workbook(caminho, data_only=True)
    mc = wb[aba_mc]
    qpc = valores_por_item(wb[aba_qpc], col_valor)
    rot = valores_por_item(wb[aba_rot], col_valor)
    atividades, avisos, item = [], [], None
    ordem = 0
    for r in range(8, mc.max_row + 1):
        a, b, c = mc.cell(r, 1).value, mc.cell(r, 2).value, mc.cell(r, 3).value
        d, e, g = mc.cell(r, 4).value, mc.cell(r, 5).value, mc.cell(r, 7).value
        cod = codigos(a)
        if cod and len(next(iter(cod)).split(".")) >= 3:
            dono = str(b or "")
            if dono and marca.lower() not in dono.lower():
                item = None
                continue
            desc = str(c or "").strip()
            if not desc and isinstance(e, str):
                desc = e
            item = {"codigo": " / ".join(sorted(cod)), "chaves": cod, "desc": desc.strip(), "n": 0}
            item["valor_qpc"], item["valor_rotula"] = achar(qpc, cod), achar(rot, cod)
            if item["valor_qpc"] is None or item["valor_rotula"] is None:
                avisos.append(f"{empresa} item {item['codigo']}: valor de contrato não encontrado")
            continue
        if item is None or a not in (None, "") or not c or not isinstance(g, (int, float)) or not isinstance(d, (int, float)):
            continue
        item["n"] += 1
        ordem += 1
        nome = str(c).strip()
        ctl = next((v for (emp, it, chave), v in CONTROLE.items()
                    if emp == empresa and it == item["codigo"] and chave in nome.lower()), None)
        atividades.append({
            "codigo": f"{empresa[:3].replace(' ', '')}-{item['codigo'].replace(' / ', '_')}-{item['n']:02d}",
            "empresa": empresa, "item": item["codigo"], "item_desc": item["desc"], "atividade": nome,
            "unidade": unidade(e), "qtd": float(d), "peso": float(g),
            "valor_qpc": round(item["valor_qpc"] or 0.0, 2), "valor_rotula": round(item["valor_rotula"] or 0.0, 2),
            "controle": ctl,
        })
    # confere pesos e totais
    por_item = {}
    for x in atividades:
        por_item.setdefault(x["item"], []).append(x)
    for it, xs in por_item.items():
        s = sum(x["peso"] for x in xs)
        if abs(s - 1) > 1e-6:
            avisos.append(f"{empresa} item {it}: soma dos pesos das atividades = {s:.4f} (esperado 1)")
    tot_qpc = sum(xs[0]["valor_qpc"] for xs in por_item.values())
    tot_rot = sum(xs[0]["valor_rotula"] for xs in por_item.values())
    return atividades, avisos, {"itens": len(por_item), "atividades": len(atividades),
                                "total_qpc": round(tot_qpc, 2), "total_rotula": round(tot_rot, 2),
                                "bm_qpc": round(sum(qpc.values()), 2), "bm_rotula": round(sum(rot.values()), 2)}


if __name__ == "__main__":
    if len(sys.argv) != 4:
        sys.exit(__doc__)
    todas, resumo = [], {}
    for emp, arq in zip(("EJ", "CMM", "GLOBO AÇOS"), sys.argv[1:4]):
        ats, avisos, res = extrair(emp, arq)
        todas += ats
        resumo[emp] = res
        print(f"\n{emp}: {res['itens']} itens, {res['atividades']} atividades | "
              f"soma dos itens QPC {res['total_qpc']:,.2f} RÓTULA {res['total_rotula']:,.2f}")
        for av in avisos:
            print("  AVISO:", av)
    with open(SAIDA, "w", encoding="utf-8") as f:
        json.dump({"fonte": "planilhas de medição (BM) de EJ, CMM e GLOBO", "atividades": todas}, f, ensure_ascii=False, indent=1)
    print(f"\nGerado: {SAIDA} ({len(todas)} atividades)")
