"""Registra no sistema o que já foi medido (e pago) em cada BM, a partir das planilhas de medição das empresas.

O BM RÓTULA mede o acumulado da produção menos o que já foi medido nos BMs anteriores. Para isso cada BM
fechado guarda, por atividade, a quantidade acumulada medida (memória de cálculo da planilha), a data de corte
e os descontos daquele BM (equipamento emprestado, faturamento direto, diesel, sinal de contrato, almoço...).

Uso:
  python registrar_bm_medido.py extrair BM_EJ_BM1.xlsx BM_CMM.xlsx saida.json [BM_EJ_BM2_rascunho.xlsx]
  python registrar_bm_medido.py carregar dados/bm_medido_bm1.json PLANO_DE_PRODUCAO.xlsm [--refazer]

O que é carregado na planilha de produção:
  BM PERÍODOS            BM1 de cada empresa com a data de corte
  BM FECHAMENTO          BM1 marcado como fechado/pago
  BM FECHAMENTO ATIVIDADES  quantidade acumulada medida em cada atividade no BM1
  BM DEDUÇÕES            descontos do BM1 (fechado) e, quando informado, os do BM2 em aberto
"""
import argparse
import datetime as dt
import json
import os
import re
import sys

import openpyxl

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(RAIZ, "ferramentas"))

# empresa -> (aba rótula, aba da memória de cálculo do BM1, coluna do valor do período na rótula)
CONFIG = {
    "EJ": ("BM1 - EJ MONTAGENS - ROTULA", "BM1-MC MONTAGEM ", "AL"),
    "CMM": ("BM - CMM - ROTULA", "BM1- MC MONTAGEM", "AM"),
}


def _codigos_atividades(empresa, caminho, aba_mc):
    """Quantidade acumulada medida (coluna N da memória de cálculo) por código de atividade do catálogo."""
    import extrair_catalogo_bm as cat
    _, _, _, _, marca = cat.CONFIG[empresa]
    mc = openpyxl.load_workbook(caminho, data_only=True)[aba_mc]
    item, saida = None, {}
    for r in range(8, mc.max_row + 1):
        a, b, c = mc.cell(r, 1).value, mc.cell(r, 2).value, mc.cell(r, 3).value
        d, g = mc.cell(r, 4).value, mc.cell(r, 7).value
        cod = cat.codigos(a)
        if cod and len(next(iter(cod)).split(".")) >= 3:
            dono = str(b or "")
            if dono and marca.lower() not in dono.lower():
                item = None
                continue
            item = {"codigo": " / ".join(sorted(cod)), "n": 0}
            continue
        if item is None or a not in (None, "") or not c or not isinstance(g, (int, float)) or not isinstance(d, (int, float)):
            continue
        item["n"] += 1
        n = mc.cell(r, 14).value
        saida[f"{empresa[:3].replace(' ', '')}-{item['codigo'].replace(' / ', '_')}-{item['n']:02d}"] = float(n) if isinstance(n, (int, float)) else 0.0
    return saida


def _descontos(caminho, aba, col):
    """Linhas de desconto entre 'TOTAL DA MEDIÇÃO DO PERIODO' e 'VALOR A FATURAR' (valor e fórmula de cada linha)."""
    ws = openpyxl.load_workbook(caminho, data_only=True)[aba]
    wf = openpyxl.load_workbook(caminho, data_only=False)[aba]
    ini = fim = None
    for r in range(1, ws.max_row + 1):
        t = str(ws.cell(r, 1).value or "").upper()
        if t.startswith("TOTAL DA MEDIÇÃO") and ini is None:
            ini = r
        if t.startswith("VALOR A FATURAR") and ini is not None:
            fim = r
            break
    medido = ws[f"{col}{ini}"].value
    fatura = ws[f"{col}{fim}"].value
    linhas = []
    for r in range(ini + 1, fim):
        desc = str(ws.cell(r, 1).value or "").strip()
        v, f = ws[f"{col}{r}"].value, wf[f"{col}{r}"].value
        if not desc or not isinstance(v, (int, float)):
            continue
        pct = None
        m = re.search(r"\(?(-?\d+(?:[.,]\d+)?)%\)?", str(f)) if isinstance(f, str) else None
        if m and "*" in str(f):
            pct = abs(float(m.group(1).replace(",", ".")))
        tipo = ("SINAL DE CONTRATO" if "SINAL" in desc.upper() and pct else
                "MEDIÇÃO ANTECIPADA" if "ANTECIPADA" in desc.upper() else
                "FATURAMENTO DIRETO" if desc.upper().startswith("FATURAMENTO DIRETO") else
                "COMBUSTÍVEL" if desc.upper().startswith("DIESEL") else
                "REFEIÇÃO" if "ALMOÇO" in desc.upper() else
                "EQUIPAMENTO EMPRESTADO" if "EQUIPAMENTO" in desc.upper() else "OUTROS")
        linhas.append({"tipo": tipo, "descricao": desc, "valor": None if pct else round(-v, 2), "percentual": pct * 1.0 if pct else None,
                       "valor_planilha": round(v, 2)})
    return {"medido": round(medido, 2), "a_faturar": round(fatura, 2), "linhas": linhas}


def extrair(ej_bm1, cmm_bm1, ej_bm2=None):
    saida = {"fonte": "planilhas de medição BM1 (EJ, CMM)", "corte": "2026-09-03", "fechado_em": "2026-09-03", "empresas": {}}
    for emp, arq in (("EJ", ej_bm1), ("CMM", cmm_bm1)):
        aba, aba_mc, col = CONFIG[emp]
        d = _descontos(arq, aba, col)
        saida["empresas"][emp] = {"bm": 1, "realizado": _codigos_atividades(emp, arq, aba_mc), **d}
    if ej_bm2:
        aba, _, col = CONFIG["EJ"]
        saida["bm2_rascunho"] = {"EJ": _descontos(ej_bm2, aba, col)}
    return saida


def carregar(dados, planilha, refazer=False):
    os.environ["OBRA198_SEM_ATUALIZAR"] = "1"
    sys.path.insert(0, os.path.join(RAIZ, "programa_obra198"))
    import excel_io as xio
    backups = os.path.join(os.path.dirname(os.path.abspath(planilha)), "backups_obra198")
    xio.garantir_abas(planilha, pasta_backup=backups)
    xio.semear_catalogo_bm(planilha, pasta_backup=backups)
    T = xio.ler_planilha(planilha)["tabelas"]
    catalogo = {}
    for a in T["bm_atividades"]:
        catalogo.setdefault(a["empresa"], []).append(a)
    celulas = []

    def limpar(tabela, linhas):
        spec = xio.TABELAS[tabela]
        for r in linhas:
            celulas.extend((spec["aba"], f"{xio.col_letra(i)}{r['linha']}", None, None) for i in range(1, len(spec["campos"]) + 1))

    def linhas_livres(tabela, reservadas):
        spec = xio.TABELAS[tabela]
        usadas = {r["linha"] for r in T[tabela]} - {r["linha"] for r in reservadas}
        n = xio.LINHA_DADOS_TABELA
        while True:
            if n not in usadas:
                yield n
            n += 1

    def escrever(tabela, registros, reservadas):
        spec, livres = xio.TABELAS[tabela], linhas_livres(tabela, reservadas)
        for campos in registros:
            lin = next(livres)
            for col, v, t in xio.valores_registro(tabela, campos):
                celulas.append((spec["aba"], f"{col}{lin}", v, t))

    # cabeçalho da coluna nova dos descontos (planilhas criadas antes do campo "% do medido")
    spec_d = xio.TABELAS["bm_deducoes"]
    celulas.append((spec_d["aba"], f"{xio.col_letra(len(spec_d['campos']))}2", spec_d["campos"][-1][1], None))

    emps = list(dados["empresas"])
    antigos = {t: [r for r in T[t] if r["empresa"] in emps and int(r.get("bm") or 0) == 1] for t in ("bm_periodos", "bm_fechamentos", "bm_fech_atividades")}
    ant_ded = [r for r in T["bm_deducoes"] if r["empresa"] in emps and (r.get("descricao") or "").startswith("[planilha]")]
    if any(antigos.values()) and not refazer:
        sys.exit("O BM1 já está registrado na planilha. Use --refazer para substituir.")
    for t, rs_ in antigos.items():
        limpar(t, rs_)
    limpar("bm_deducoes", ant_ded)

    novos = {"bm_periodos": [], "bm_fechamentos": [], "bm_fech_atividades": [], "bm_deducoes": []}
    for emp, d in dados["empresas"].items():
        novos["bm_periodos"].append({"empresa": emp, "bm": 1, "corte": dados["corte"], "obs": "BM1 pago (planilha de medição)"})
        novos["bm_fechamentos"].append({"empresa": emp, "bm": 1, "fechado_em": dados["fechado_em"], "equip": 0, "comb": 0, "litros": 0,
                                        "obs": f"BM1 já medido e pago: R$ {d['medido']:,.2f} medido, R$ {d['a_faturar']:,.2f} faturado (descontos lançados em BM DEDUÇÕES)"})
        for a in catalogo.get(emp, []):
            novos["bm_fech_atividades"].append({"empresa": emp, "bm": 1, "codigo": a["codigo"], "item": a["item"],
                                                "realizado": d["realizado"].get(a["codigo"], 0.0), "qtd": a["qtd"], "peso": a["peso"],
                                                "valor_qpc": a.get("valor_qpc"), "valor_rotula": a.get("valor_rotula")})
        for l in d["linhas"]:
            novos["bm_deducoes"].append({"empresa": emp, "bm": 1, "data": dados["corte"], "tipo": l["tipo"], "valor": l["valor"],
                                         "percentual": l["percentual"], "contrato": "RÓTULA", "descricao": "[planilha] " + l["descricao"]})
    for emp, d in (dados.get("bm2_rascunho") or {}).items():
        for l in d["linhas"]:
            novos["bm_deducoes"].append({"empresa": emp, "bm": 2, "data": dt.date.today().isoformat(), "tipo": l["tipo"], "valor": l["valor"],
                                         "percentual": l["percentual"], "contrato": "RÓTULA",
                                         "descricao": "[planilha] " + l["descricao"] + " (rascunho do BM2: conferir)"})
    for tabela, regs in novos.items():
        escrever(tabela, regs, antigos.get(tabela, ant_ded if tabela == "bm_deducoes" else []))
    xio.gravar_celulas(planilha, celulas, pasta_backup=backups)
    return {t: len(r) for t, r in novos.items()}


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)
    e = sub.add_parser("extrair")
    e.add_argument("ej_bm1")
    e.add_argument("cmm_bm1")
    e.add_argument("saida")
    e.add_argument("ej_bm2", nargs="?")
    c = sub.add_parser("carregar")
    c.add_argument("json")
    c.add_argument("planilha")
    c.add_argument("--refazer", action="store_true")
    a = ap.parse_args()
    if a.cmd == "extrair":
        d = extrair(a.ej_bm1, a.cmm_bm1, a.ej_bm2)
        with open(a.saida, "w", encoding="utf-8") as f:
            json.dump(d, f, ensure_ascii=False, indent=1)
        for emp, x in d["empresas"].items():
            print(f"{emp}: medido R$ {x['medido']:,.2f} · a faturar R$ {x['a_faturar']:,.2f} · {len(x['linhas'])} descontos · "
                  f"{sum(1 for v in x['realizado'].values() if v)} atividades com quantidade")
    else:
        print("Gravado:", json.dumps(carregar(json.load(open(a.json, encoding="utf-8")), a.planilha, a.refazer)))


if __name__ == "__main__":
    main()
