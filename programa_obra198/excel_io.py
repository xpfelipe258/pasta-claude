"""Leitura e gravação da planilha de produção.

Leitura: openpyxl (somente leitura, não altera o arquivo).
Gravação: edição cirúrgica do XML das células alteradas dentro do .xlsm,
preservando macros, gráficos, formatação e todas as demais abas.
"""
import collections
import datetime as dt
import json
import os
import re
import shutil
import tempfile
import zipfile

import openpyxl
from lxml import etree

BASE = os.path.dirname(os.path.abspath(__file__))
RAIZ = os.path.dirname(BASE)
CAMINHO_INVENTARIO_IFC = os.path.join(RAIZ, "dados", "ifc_r0d_inventario.json")
CAMINHO_REGRAS_BAIXA = os.path.join(BASE, "regras_baixa.json")

NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
NS_REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
NS_PKG_REL = "http://schemas.openxmlformats.org/package/2006/relationships"

ABA_CONSOLIDADO = "CONTROLE PRODUÇÃO POR EMPRESAS"
ABA_METAS = "METAS EMPRESAS"
ABA_CLIENTE = "METAS CLIENTE"
ABA_ESCOPO = "ESCOPO POR EMPRESA"
ABA_IMPACTO = "LISTA DE IMPACTO"
ABA_PAINEL = "PAINEL KPI"
ABAS_IGNORADAS = {ABA_CONSOLIDADO, "CONTROLE KPI"}

LINHA_INICIO_DADOS = 8
LINHA_FIM_DADOS = 364
LINHA_INICIO_METAS = 17
LINHAS_IMPACTO = range(5, 35)
EXCEL_EPOCH = dt.date(1899, 12, 30)


class PlanilhaBloqueada(Exception):
    pass


class CelulaProtegida(Exception):
    pass


# ---------------------------------------------------------------- utilidades

def col_letra(n):
    s = ""
    while n:
        n, r = divmod(n - 1, 26)
        s = chr(65 + r) + s
    return s


def col_numero(letras):
    n = 0
    for ch in letras:
        n = n * 26 + ord(ch) - 64
    return n


def separa_ref(ref):
    m = re.match(r"([A-Z]+)(\d+)$", ref)
    return m.group(1), int(m.group(2))


def iso(v):
    if isinstance(v, dt.datetime):
        return v.date().isoformat()
    if isinstance(v, dt.date):
        return v.isoformat()
    return None


def numero(v):
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return v
    if isinstance(v, str):
        t = v.strip().replace(",", ".")
        try:
            return float(t) if t else None
        except ValueError:
            return None
    return None


def eh_formula(v):
    return isinstance(v, str) and v.startswith("=")


# ------------------------------------------------------------------- leitura

def _linhas(ws, max_col=None):
    if hasattr(ws, "reset_dimensions"):
        ws.reset_dimensions()
    return [list(r) for r in ws.iter_rows(values_only=True, max_col=max_col)]


def _cel(linhas, linha, col):
    """linha/col 1-based; devolve None fora dos limites."""
    if linha - 1 >= len(linhas):
        return None
    r = linhas[linha - 1]
    return r[col - 1] if col - 1 < len(r) else None


def _avaliar_escopo(escopo, cliente, ref):
    """Avalia a célula de ESCOPO POR EMPRESA sem depender do cache do Excel."""
    m = re.match(r"([A-Z]+)(\d+)$", ref)
    if not m:
        return None
    col, lin = m.group(1), int(m.group(2))

    def qtd_cliente(l):
        return numero(_cel(cliente, l, 2))

    def valor(c, l, prof=0):
        if prof > 5:
            return None
        v = _cel(escopo, l, col_numero(c))
        if not eh_formula(v):
            return numero(v)
        f = v.replace(" ", "")
        if "'METASCLIENTE'!B" in f.replace(" ", ""):
            return qtd_cliente(l)
        if "ROUNDUP" in f:
            d = valor("D", l, prof + 1)
            return None if d is None else -(-d // 2)
        if "-N(E" in f:
            d = valor("D", l, prof + 1)
            e = valor("E", l, prof + 1) or 0
            return None if d is None else d - e
        mm = re.match(r"=([A-Z]+)(\d+)$", f)
        if mm:
            return valor(mm.group(1), int(mm.group(2)), prof + 1)
        return None

    return valor(col, lin)


def ler_planilha(caminho):
    wb = openpyxl.load_workbook(caminho, read_only=True, keep_vba=False)
    try:
        return _montar_modelo(wb)
    finally:
        wb.close()


def carregar_inventario_ifc():
    if not os.path.exists(CAMINHO_INVENTARIO_IFC):
        return None
    with open(CAMINHO_INVENTARIO_IFC, encoding="utf-8") as f:
        return json.load(f)


def carregar_regras_baixa():
    if not os.path.exists(CAMINHO_REGRAS_BAIXA):
        return None
    with open(CAMINHO_REGRAS_BAIXA, encoding="utf-8") as f:
        return json.load(f)


def calcular_consumo_servico(servico, delta, regras):
    """Calcula o consumo de materiais para uma dada produção.

    Retorna lista de dicts [{codigo, descricao, quantidade}] ou lista vazia.
    """
    if not regras or delta <= 0:
        return []
    mapa = regras.get("consumo_por_servico", {})
    conf = mapa.get(servico)
    if not conf:
        return []
    resultado = []
    if "itens" in conf:
        for item in conf["itens"]:
            q = round(item["por_unidade"] * delta, 2)
            if q > 0:
                resultado.append({"codigo": item["codigo"], "descricao": item.get("descricao", ""),
                                  "quantidade": q})
    return resultado


def _consumo_evento(ev, regras):
    """Devolve dict {codigo_material: quantidade} consumido por um evento."""
    tipo = (ev.get("tipo") or "").upper()
    letra = (ev.get("letra") or "").upper()
    faixa = (ev.get("faixa") or "").upper()
    consumo = collections.Counter()

    if tipo.startswith("VIGA_") and tipo.endswith("_MONTADA"):
        for e in regras.get("eventos", []):
            if e["codigo"] == tipo:
                for b in e.get("baixa", []):
                    consumo[b["codigo"]] += b["quantidade"]
                break
        return consumo

    if tipo == "JOIST_ICADA":
        n = ev.get("n_joists") or 0
        try:
            n = int(n)
        except (TypeError, ValueError):
            n = 0
        if n <= 0:
            return consumo
        # cada joist tem 4 parafusos em cada uma das 2 pontas (eixo ini e eixo fim), do
        # tipo definido pela LETRA do par de vigas (A/B/BC/C/CD/D/DE/E/F/FG/G/H).
        # Se só a faixa veio, usamos a primeira letra da faixa (borda apoio) como fallback.
        crit = regras.get("criterio_por_letra", {}) or {}
        chave_letra = letra if letra in crit else (
            faixa if faixa in crit else (faixa[:1] if faixa and faixa[:1] in crit else ""))
        par = crit.get(chave_letra) or {}
        parafuso = par.get("parafuso")
        for e in regras.get("eventos", []):
            if e["codigo"] != "JOIST_ICADA":
                continue
            for b in e.get("baixa_por_joist", []):
                cod = b.get("codigo")
                q = b.get("quantidade", 0) * n
                if cod == "PARAFUSO_FIX_JOIST" and parafuso:
                    consumo[parafuso] += q
                elif cod and cod != "PARAFUSO_FIX_JOIST":
                    consumo[cod] += q
        return consumo

    return consumo


def eventos_de_apontamentos(apontamentos, regras):
    """Converte as linhas de APONTAMENTO MONTAGEM nos eventos que o motor de baixa entende."""
    crit = (regras or {}).get("criterio_por_letra", {})
    saida = []
    for a in apontamentos or []:
        tipo = (a.get("tipo") or "").upper()
        base = {"data": a.get("data"), "empresa": a.get("empresa"), "letra": a.get("letra")}
        if tipo == "JOIST":
            saida.append({**base, "tipo": "JOIST_ICADA", "rua": a.get("rua"), "faixa": a.get("faixa"),
                          "n_joists": a.get("qtd")})
        elif tipo == "VIGA":
            ev = (crit.get((a.get("letra") or "").upper()) or {}).get("evento_viga")
            if ev:
                saida.append({**base, "tipo": ev, "eixo": a.get("eixo"), "viga_id": a.get("viga_id")})
    return saida


def resumir_estoque_ifc(eventos, remessas, inventarios, regras, inv_ifc):
    """Consolida consumo teórico, saldo, cobertura, perda e acuracidade por material."""
    if not regras:
        return None

    # 1) consumo teórico acumulado por material
    consumo = collections.Counter()
    consumo_por_dia = collections.defaultdict(collections.Counter)
    joists_por_rua_faixa = collections.defaultdict(int)
    alertas = []
    limite = (regras.get("limites_ifc") or {}).get("joists_por_rua_faixa_maximo") or 6

    letra_para_faixa = regras.get("letra_para_faixa_do_par") or {}
    for ev in eventos or []:
        c = _consumo_evento(ev, regras)
        consumo.update(c)
        if ev.get("data"):
            consumo_por_dia[ev["data"]].update(c)
        if (ev.get("tipo") or "").upper() == "JOIST_ICADA":
            n = int(ev.get("n_joists") or 0)
            faixa = (ev.get("faixa") or letra_para_faixa.get((ev.get("letra") or "").upper()) or "").upper()
            chave = f"{ev.get('rua', '')}/{faixa}"
            joists_por_rua_faixa[chave] += n
            if joists_por_rua_faixa[chave] > limite:
                alertas.append({
                    "tipo": "joists_excedeu_ifc",
                    "chave": chave,
                    "mensagem": f"Rua/faixa {chave}: {joists_por_rua_faixa[chave]} joists apontadas (máx. IFC {limite}).",
                })

    # 2) chegou (soma remessas) e inventário mais recente por material
    chegou = collections.Counter()
    for r in remessas or []:
        q = r.get("quantidade")
        cod = r.get("codigo")
        if q and cod:
            chegou[cod] += float(q)

    ultimo_inv = {}
    for i in inventarios or []:
        cod = i.get("codigo")
        q = i.get("quantidade")
        d = i.get("data")
        if cod and q is not None and d and (cod not in ultimo_inv or d >= ultimo_inv[cod]["data"]):
            ultimo_inv[cod] = {"data": d, "quantidade": float(q)}

    # 3) por material: saldo teórico = chegou − consumo; físico − teórico = perda; acuracidade
    codigos = set(consumo) | set(chegou) | set(ultimo_inv)
    materiais = []
    for cod in sorted(codigos):
        c = float(consumo.get(cod, 0))
        entrou = float(chegou.get(cod, 0))
        saldo_teo = entrou - c
        fis = ultimo_inv.get(cod, {}).get("quantidade")
        perda = None
        acuracia = None
        if fis is not None:
            perda = fis - saldo_teo
            f, s = max(fis, 0), max(saldo_teo, 0)
            acuracia = (min(f, s) / max(f, s)) if max(f, s) else 1.0
        materiais.append({
            "codigo": cod, "consumo_teorico": c, "chegou": entrou,
            "saldo_teorico": saldo_teo,
            "fisico": fis, "data_inventario": (ultimo_inv.get(cod) or {}).get("data"),
            "perda": perda, "acuracidade": acuracia,
        })

    # 4) consumo por rua/faixa vs. projetado (limite IFC)
    joists_proj = ((inv_ifc or {}).get("joists_projetadas") or {}).get("por_rua_faixa", {})
    ruas = []
    for rua, faixas in joists_proj.items():
        for fx, proj in faixas.items():
            apont = joists_por_rua_faixa.get(f"{rua}/{fx}", 0)
            ruas.append({"rua": rua, "faixa": fx, "projetado": proj, "apontado": apont,
                         "saldo": proj - apont})

    # 5) cobertura (dias) por material: quantos dias o saldo teórico dura ao ritmo médio
    hoje = dt.date.today()
    for m in materiais:
        m["cobertura_dias"] = None
        c = m["consumo_teorico"]
        if not c:
            continue
        # ritmo médio: consumo dividido pelo intervalo de dias entre 1º e último evento
        datas_c = sorted(consumo_por_dia)
        if len(datas_c) >= 2:
            try:
                dias = (dt.date.fromisoformat(datas_c[-1]) - dt.date.fromisoformat(datas_c[0])).days or 1
                ritmo = c / dias  # unid/dia
                if ritmo > 0 and m["saldo_teorico"] is not None:
                    m["cobertura_dias"] = round(m["saldo_teorico"] / ritmo, 1)
            except ValueError:
                pass

    return {
        "materiais": materiais,
        "por_rua_faixa": ruas,
        "consumo_por_dia": {d: dict(c) for d, c in consumo_por_dia.items()},
        "alertas": alertas,
        "gerado_em": hoje.isoformat(),
    }


def _montar_modelo(wb):
    cons = _linhas(wb[ABA_CONSOLIDADO], max_col=1)
    datas = {}
    for i in range(6, len(cons) + 1):
        d = iso(_cel(cons, i, 1))
        if d:
            datas[i + 2] = d  # linha na aba da empresa = linha do consolidado + 2

    escopo = _linhas(wb[ABA_ESCOPO], max_col=10)
    cliente = _linhas(wb[ABA_CLIENTE], max_col=14)

    empresas = []
    for ws in wb.worksheets:
        if not ws.title.startswith("CONTROLE ") or ws.title in ABAS_IGNORADAS:
            continue
        linhas = _linhas(ws)
        cab = linhas[6] if len(linhas) > 6 else []
        servicos, col_obs, frente = [], None, ""
        for c in range(2, len(cab) + 1):
            nome = _cel(linhas, 7, c)
            if nome is None:
                continue
            if str(nome).upper().startswith("OBSERVA"):
                col_obs = c
                break
            frente = _cel(linhas, 2, c) or frente
            ref_escopo = _cel(linhas, 4, c)
            qtd = None
            m = re.search(r"'ESCOPO POR EMPRESA'!([A-Z]+\d+)", str(ref_escopo or ""))
            if m:
                qtd = _avaliar_escopo(escopo, cliente, m.group(1))
            elif not eh_formula(ref_escopo):
                qtd = numero(ref_escopo)
            servicos.append({
                "col": col_letra(c),
                "nome": str(nome).strip(),
                "frente": str(frente).strip(),
                "id": str(_cel(linhas, 3, c) or "").strip(),
                "escopo": qtd,
            })

        registros = {}
        for lin, data in datas.items():
            if lin > LINHA_FIM_DADOS:
                continue
            vals = {}
            for s in servicos:
                v = _cel(linhas, lin, col_numero(s["col"]))
                if v is not None and v != "":
                    vals[s["col"]] = v if isinstance(v, (int, float)) else str(v)
            obs = _cel(linhas, lin, col_obs) if col_obs else None
            if vals or obs:
                reg = {"v": vals}
                if obs:
                    reg["obs"] = str(obs)
                registros[data] = reg

        empresas.append({
            "aba": ws.title,
            "nome": ws.title[len("CONTROLE "):],
            "servicos": servicos,
            "col_obs": col_letra(col_obs) if col_obs else None,
            "registros": registros,
        })

    tabelas = ler_tabelas(wb)
    inv_ifc = carregar_inventario_ifc()
    regras = carregar_regras_baixa()
    eventos = list(tabelas.get("estoque_eventos") or []) + eventos_de_apontamentos(tabelas.get("apontamentos"), regras)
    resumo = resumir_estoque_ifc(
        eventos, tabelas.get("estoque_remessas"),
        tabelas.get("estoque_inventario"), regras, inv_ifc) if regras else None
    return {
        "datas": sorted(datas.values()),
        "linhas_por_data": {d: l for l, d in datas.items()},
        "empresas": empresas,
        "metas": _ler_metas(wb[ABA_METAS]),
        "cliente": _ler_cliente(cliente),
        "impactos": _ler_impactos(wb[ABA_IMPACTO]),
        "referencia": iso(wb[ABA_PAINEL]["B3"].value) if ABA_PAINEL in wb.sheetnames else None,
        "tabelas": tabelas,
        "abas_mensais_equip": [n for n in wb.sheetnames if n.upper().startswith("USO DE EQUIPAMENTO")],
        "estoque_ifc": {"inventario": inv_ifc, "regras": regras, "resumo": resumo},
    }


def _ler_metas(ws):
    linhas = _linhas(ws, max_col=14)
    metas = []
    for lin in range(LINHA_INICIO_METAS, len(linhas) + 1):
        emp = _cel(linhas, lin, 1)
        ini = iso(_cel(linhas, lin, 3))
        if not emp or not ini:
            continue
        real = _cel(linhas, lin, 12)
        metas.append({
            "linha": lin,
            "empresa": str(emp).strip(),
            "servico": str(_cel(linhas, lin, 2) or "").strip(),
            "inicio": ini,
            "fim": iso(_cel(linhas, lin, 4)) or ini,
            "corte": iso(_cel(linhas, lin, 5)),
            "meta_dia": numero(_cel(linhas, lin, 6)) or 0,
            "du": numero(_cel(linhas, lin, 7)) or 0,
            "gap": numero(_cel(linhas, lin, 9)) or 0,
            "realizado_manual": not eh_formula(real),
        })
    return metas


def _ler_cliente(linhas):
    itens = []
    for lin in range(5, len(linhas) + 1):
        serv = _cel(linhas, lin, 1)
        if not serv or eh_formula(serv):
            continue
        itens.append({
            "linha": lin,
            "servico": str(serv).strip(),
            "qtd": numero(_cel(linhas, lin, 2)),
            "inicio_plan": iso(_cel(linhas, lin, 3)),
            "meta_dia": numero(_cel(linhas, lin, 5)),
            "du_semana": numero(_cel(linhas, lin, 6)),
            "responsavel": _cel(linhas, lin, 9),
            "obs": _cel(linhas, lin, 10),
            "prazo": iso(_cel(linhas, lin, 13)),
            "peso": numero(_cel(linhas, lin, 12)) or 1,
            "frente": None if eh_formula(_cel(linhas, lin, 14)) else _cel(linhas, lin, 14),
        })
    return itens


def _ler_impactos(ws):
    linhas = _linhas(ws, max_col=8)
    itens = []
    for lin in LINHAS_IMPACTO:
        vals = [_cel(linhas, lin, c) for c in range(1, 9)]
        if not any(v not in (None, "") for v in vals[1:]):
            continue
        itens.append({
            "linha": lin,
            "data": iso(vals[1]) or (str(vals[1]) if vals[1] else None),
            "servico": vals[2],
            "motivo": vals[3],
            "solucionado": iso(vals[4]) or (str(vals[4]) if vals[4] else None),
            "tempo": vals[5],
            "quando": iso(vals[6]) or (str(vals[6]).strip() if vals[6] else None),
            "paralisacao": vals[7],
        })
    return itens


def primeira_linha_impacto_livre(caminho):
    wb = openpyxl.load_workbook(caminho, read_only=True)
    try:
        linhas = _linhas(wb[ABA_IMPACTO], max_col=8)
    finally:
        wb.close()
    for lin in LINHAS_IMPACTO:
        if all(_cel(linhas, lin, c) in (None, "") for c in range(2, 9)):
            return lin
    return None


# --------------------------------------------------------- tabelas do sistema
# Abas criadas pelo programa (título na linha 1, cabeçalho na linha 2, dados a partir da 3).

TABELAS = {
    "materiais": {
        "aba": "ESTOQUE MATERIAIS",
        "titulo": "ESTOQUE DE MATERIAIS — cadastro. Consumo por unidade = quanto do material é gasto para executar 1 unidade do serviço vinculado.",
        "cor": "FF7F6000",
        "campos": [
            ("codigo", "Código", "txt", 12), ("material", "Material", "txt", 34), ("unidade", "Unidade", "txt", 10),
            ("servico", "Serviço vinculado", "txt", 22), ("empresa", "Empresa (vazio = todas)", "txt", 16),
            ("coef", "Consumo por unidade de serviço", "num", 14), ("saldo_inicial", "Saldo inicial", "num", 12),
            ("data_saldo", "Data do saldo inicial", "data", 14), ("minimo", "Estoque mínimo", "num", 12),
            ("fornecedor", "Fornecedor", "txt", 20), ("prazo_reposicao", "Prazo reposição (dias)", "num", 12),
            ("obs", "Observação", "txt", 30),
        ],
    },
    "movimentos": {
        "aba": "MOVIMENTAÇÃO ESTOQUE",
        "titulo": "MOVIMENTAÇÃO DE ESTOQUE — ENTRADA (recebimento), SAÍDA (perda, transferência, consumo não apropriado) ou AJUSTE (+/−, inventário).",
        "cor": "FF7F6000",
        "campos": [
            ("data", "Data", "data", 12), ("codigo", "Código material", "txt", 14), ("tipo", "Tipo", "txt", 10),
            ("quantidade", "Quantidade", "num", 12), ("documento", "Documento (NF / romaneio)", "txt", 20),
            ("empresa", "Empresa", "txt", 14), ("obs", "Observação", "txt", 34),
        ],
    },
    "equipamentos": {
        "aba": "CADASTRO EQUIPAMENTOS",
        "titulo": "CADASTRO DE EQUIPAMENTOS — valor unitário conforme a forma de cobrança (diária, hora, turno ou mensal).",
        "cor": "FF375623",
        "campos": [
            ("equipamento", "Equipamento", "txt", 22), ("tipo", "Tipo", "txt", 16), ("locadora", "Locadora / proprietário", "txt", 20),
            ("cobranca", "Cobrança", "txt", 12), ("valor", "Valor unitário (R$)", "num", 14),
            ("consumo_lh", "Consumo estimado (L/h)", "num", 14), ("situacao", "Situação", "txt", 14), ("obs", "Observação", "txt", 30),
        ],
    },
    "usos": {
        "aba": "USO EQUIPAMENTOS",
        "titulo": "USO DE EQUIPAMENTOS E ABASTECIMENTO — uma linha por uso ou abastecimento. Empresa compartilhada: EJ/CMM (custo dividido igualmente).",
        "cor": "FF375623",
        "campos": [
            ("data", "Data", "data", 12), ("equipamento", "Equipamento", "txt", 18), ("empresa", "Empresa", "txt", 14),
            ("uso", "Uso", "txt", 14), ("quantidade", "Qtd (diárias/horas)", "num", 12), ("custo", "Custo equipamento (R$)", "num", 14),
            ("litros", "Combustível (L)", "num", 12), ("preco_litro", "Preço litro (R$)", "num", 12),
            ("custo_combustivel", "Custo combustível (R$)", "num", 14), ("operador", "Operador", "txt", 14),
            ("obs", "Observação", "txt", 28), ("origem", "Origem", "txt", 30),
        ],
    },
    "bm_atividades": {
        "aba": "BM ATIVIDADES",
        "titulo": "BM — CATÁLOGO DE ATIVIDADES POR EMPRESA (memória de cálculo). Cada atividade tem peso dentro do item e quantidade contratada; o valor do item é o do contrato QPC e RÓTULA. 'Coluna no controle' liga a atividade a um serviço do lançamento de produção.",
        "cor": "FF1F4E79",
        "campos": [
            ("codigo", "Código", "txt", 16), ("empresa", "Empresa", "txt", 14), ("item", "Item do BM", "txt", 12),
            ("item_desc", "Descrição do item", "txt", 40), ("atividade", "Atividade (serviço)", "txt", 44), ("unidade", "Unidade", "txt", 9),
            ("qtd", "Quantidade contratada", "num", 14), ("peso", "Peso da atividade no item", "num", 12),
            ("valor_qpc", "Valor do item QPC (R$)", "num", 16), ("valor_rotula", "Valor do item RÓTULA (R$)", "num", 16),
            ("controle", "Coluna no controle de produção", "txt", 20), ("obs", "Observação", "txt", 30),
        ],
    },
    "bm_periodos": {
        "aba": "BM PERÍODOS",
        "titulo": "BM — DATA DE CORTE DE CADA MEDIÇÃO. O período vai do dia seguinte ao corte anterior até a data de corte (inclusive).",
        "cor": "FF1F4E79",
        "campos": [
            ("empresa", "Empresa", "txt", 14), ("bm", "Nº do BM", "num", 10), ("corte", "Data de corte", "data", 14),
            ("obs", "Observação", "txt", 40),
        ],
    },
    "bm_apontamentos": {
        "aba": "BM APONTAMENTO",
        "titulo": "BM — APONTAMENTO POR SERVIÇO. Quantidade executada por dia em cada atividade do catálogo (as atividades ligadas a uma coluna do controle de produção são lançadas lá).",
        "cor": "FF1F4E79",
        "campos": [
            ("data", "Data", "data", 12), ("empresa", "Empresa", "txt", 14), ("codigo", "Código da atividade", "txt", 16),
            ("quantidade", "Quantidade", "num", 12), ("obs", "Observação / nº RDO", "txt", 34),
        ],
    },
    "bm_deducoes": {
        "aba": "BM DEDUÇÕES",
        "titulo": "BM — DESCONTOS DO BM (equipamento emprestado, faturamento direto, diesel, sinal de contrato, medição antecipada, almoço...), como nas planilhas de medição. Valor negativo = crédito. Com % preenchido, o desconto é esse percentual do medido no período. Equipamentos lançados em USO EQUIPAMENTOS também são descontados automaticamente.",
        "cor": "FF1F4E79",
        "campos": [
            ("empresa", "Empresa", "txt", 14), ("bm", "Nº do BM", "num", 10), ("data", "Data", "data", 12),
            ("tipo", "Tipo", "txt", 16), ("descricao", "Descrição", "txt", 40), ("valor", "Valor a deduzir (R$)", "num", 16),
            ("contrato", "Contrato (RÓTULA, QPC ou AMBOS)", "txt", 18),
            ("percentual", "% do medido no período (ex.: 10 = sinal de contrato; vazio = usa o valor)", "num", 18),
            ("item_uso", "Item / equipamento (ex.: Guindaste Sany)", "txt", 26), ("qtd", "Qtd de uso", "num", 10),
            ("unidade", "Unidade (diária, hora...)", "txt", 12), ("valor_unit", "Valor unitário (R$); total = qtd x valor unitário", "num", 16),
        ],
    },
    "bm_fechamentos": {
        "aba": "BM FECHAMENTO",
        "titulo": "BM — FECHAMENTOS (não edite). Foto do BM no momento do fechamento: equipamento e combustível descontados. Use Medição (BM) > Fechar BM / Reabrir.",
        "cor": "FF7F0000",
        "campos": [
            ("empresa", "Empresa", "txt", 14), ("bm", "Nº do BM", "num", 10), ("fechado_em", "Fechado em", "data", 14),
            ("equip", "Equipamentos (R$)", "num", 16), ("comb", "Combustível (R$)", "num", 16), ("litros", "Litros", "num", 10),
            ("obs", "Observação", "txt", 30),
        ],
    },
    "bm_fech_atividades": {
        "aba": "BM FECHAMENTO ATIVIDADES",
        "titulo": "BM — FOTO DAS ATIVIDADES NO FECHAMENTO (não edite). Realizado acumulado, quantidade, peso e valor do item congelados em cada BM fechado.",
        "cor": "FF7F0000",
        "campos": [
            ("empresa", "Empresa", "txt", 14), ("bm", "Nº do BM", "num", 10), ("codigo", "Código da atividade", "txt", 16),
            ("item", "Item", "txt", 10), ("realizado", "Realizado acumulado", "num", 14), ("qtd", "Quantidade contratada", "num", 14),
            ("peso", "Peso no item", "num", 10), ("valor_qpc", "Valor do item QPC (R$)", "num", 16),
            ("valor_rotula", "Valor do item RÓTULA (R$)", "num", 16),
        ],
    },
    "estoque_eventos": {
        "aba": "ESTOQUE EVENTOS",
        "titulo": "ESTOQUE — EVENTOS DE PRODUÇÃO que disparam baixa de material. Tipos: VIGA_APOIO_MONTADA, VIGA_INTERM_MONTADA (identifique a viga por eixo/letra) ou JOIST_ICADA (informe rua, faixa e nº de joists içadas naquele par). O material sai automaticamente conforme regras_baixa.json.",
        "cor": "FF7F6000",
        "campos": [
            ("data", "Data", "data", 12), ("tipo", "Tipo de evento", "txt", 20),
            ("empresa", "Empresa", "txt", 14), ("eixo", "Eixo (viga) ou rua (joist)", "txt", 12),
            ("letra", "Letra do eixo / faixa", "txt", 8), ("rua", "Rua (par de eixos, ex.: 10-11)", "txt", 10),
            ("faixa", "Faixa (AB..GH)", "txt", 6), ("n_joists", "Qtd de joists (0..6)", "num", 8),
            ("viga_id", "ID da viga (opcional)", "txt", 14), ("obs", "Observação / nº RDO", "txt", 34),
        ],
    },
    "apontamentos": {
        "aba": "APONTAMENTO MONTAGEM",
        "titulo": "APONTAMENTO DE MONTAGEM (não edite à mão). Joists por rua/faixa/viga de apoio e vigas por eixo/letra, gravados pela tela Apontamento de montagem. Cada linha soma na coluna IÇAMENTO JOIST ou VIGAS do controle de produção (base do BM e do cronograma) e define a baixa de fixadores.",
        "cor": "FF7F6000",
        "campos": [
            ("data", "Data", "data", 12), ("empresa", "Empresa", "txt", 14), ("tipo", "Tipo (JOIST ou VIGA)", "txt", 12),
            ("rua", "Rua (par de eixos, ex.: 11-12)", "txt", 12), ("faixa", "Faixa (AB..GH)", "txt", 8),
            ("eixo", "Eixo da viga", "txt", 8), ("letra", "Letra da viga (posição)", "txt", 8),
            ("viga_id", "ID da viga", "txt", 16), ("qtd", "Quantidade", "num", 10),
            ("lanca_producao", "Somou na produção? (SIM/NÃO)", "txt", 14), ("obs", "Observação / nº RDO", "txt", 34),
            ("slot", "Nº da joist na rua (1 a 43)", "num", 12),
        ],
    },
    "estoque_remessas": {
        "aba": "ESTOQUE REMESSAS",
        "titulo": "ESTOQUE — REMESSAS que chegaram na obra. Usada como base do 'chegou' (entrada) do material. Uma linha por remessa/material.",
        "cor": "FF7F6000",
        "campos": [
            ("data", "Data", "data", 12), ("codigo", "Código material", "txt", 14),
            ("descricao", "Descrição", "txt", 34), ("quantidade", "Quantidade", "num", 12),
            ("documento", "Nº NF / romaneio", "txt", 20), ("fornecedor", "Fornecedor", "txt", 20),
            ("obs", "Observação", "txt", 28),
        ],
    },
    "estoque_inventario": {
        "aba": "ESTOQUE INVENTÁRIO",
        "titulo": "ESTOQUE — CONTAGEM FÍSICA. Comparado ao saldo teórico calculado por eventos × regras (base do IFC) para gerar perda e acuracidade.",
        "cor": "FF7F6000",
        "campos": [
            ("data", "Data", "data", 12), ("codigo", "Código material", "txt", 14),
            ("quantidade", "Quantidade contada", "num", 12), ("responsavel", "Responsável", "txt", 20),
            ("obs", "Observação", "txt", 34),
        ],
    },
}
LINHA_DADOS_TABELA = 3

INICIO_DADOS_ABA = {ABA_IMPACTO: 5, ABA_METAS: LINHA_INICIO_METAS}
INICIO_DADOS_ABA.update({t["aba"]: LINHA_DADOS_TABELA for t in TABELAS.values()})


def _ler_tabela(wb, spec):
    if spec["aba"] not in wb.sheetnames:
        return []
    campos = spec["campos"]
    linhas = _linhas(wb[spec["aba"]], max_col=len(campos))
    itens = []
    for lin in range(LINHA_DADOS_TABELA, len(linhas) + 1):
        vals = [_cel(linhas, lin, c) for c in range(1, len(campos) + 1)]
        if all(v in (None, "") for v in vals):
            continue
        item = {"linha": lin}
        for (chave, _, tipo, _), v in zip(campos, vals):
            if tipo == "data":
                item[chave] = iso(v) or (str(v) if v not in (None, "") else None)
            elif tipo == "num":
                item[chave] = numero(v)
            else:
                item[chave] = None if v in (None, "") else str(v).strip()
        itens.append(item)
    return itens


def ler_tabelas(wb):
    return {nome: _ler_tabela(wb, spec) for nome, spec in TABELAS.items()}


def valores_registro(tabela, campos):
    """Converte o dicionário vindo da tela em células (col, valor, tipo) da tabela."""
    spec = TABELAS[tabela]
    saida = []
    for i, (chave, _, tipo, _) in enumerate(spec["campos"], start=1):
        if chave not in campos:
            continue
        v = campos[chave]
        if isinstance(v, str):
            v = v.strip()
        if v in ("", None):
            v = None
        elif tipo == "num":
            n = numero(v)
            if n is None:
                raise ValueError(f"Valor numérico inválido em '{chave}': {v}")
            v = int(n) if float(n).is_integer() else n
        elif tipo == "data":
            dt.date.fromisoformat(v)
        saida.append((col_letra(i), v, "data" if tipo == "data" and v else None))
    return saida


# ------------------------------------------------------------------ gravação

def _mapa_abas(zf):
    wb = etree.fromstring(zf.read("xl/workbook.xml"))
    rels = etree.fromstring(zf.read("xl/_rels/workbook.xml.rels"))
    alvo = {}
    for r in rels.findall(f"{{{NS_PKG_REL}}}Relationship"):
        t = r.get("Target")
        t = t.lstrip("/") if t.startswith("/") else "xl/" + t
        alvo[r.get("Id")] = t
    return {s.get("name"): alvo[s.get(f"{{{NS_REL}}}id")]
            for s in wb.iter(f"{{{NS}}}sheet")}


def _para_serial(data_iso):
    d = dt.date.fromisoformat(data_iso)
    return (d - EXCEL_EPOCH).days


def _obter_linha(sheet_data, n):
    for row in sheet_data.findall(f"{{{NS}}}row"):
        rn = int(row.get("r"))
        if rn == n:
            return row
        if rn > n:
            novo = etree.Element(f"{{{NS}}}row", r=str(n))
            row.addprevious(novo)
            return novo
    novo = etree.SubElement(sheet_data, f"{{{NS}}}row", r=str(n))
    return novo


def _estilo_coluna(sheet_data, col, inicio):
    """Estilo de uma célula de dados existente na mesma coluna, para células novas."""
    for c in sheet_data.iter(f"{{{NS}}}c"):
        letras, lin = separa_ref(c.get("r"))
        if letras == col and lin >= inicio and c.get("s"):
            return c.get("s")
    return None


def _obter_celula(sheet_data, ref, inicio):
    col, lin = separa_ref(ref)
    row = _obter_linha(sheet_data, lin)
    alvo = col_numero(col)
    for c in row.findall(f"{{{NS}}}c"):
        cn = col_numero(separa_ref(c.get("r"))[0])
        if cn == alvo:
            return c
        if cn > alvo:
            novo = etree.Element(f"{{{NS}}}c", r=ref)
            c.addprevious(novo)
            break
    else:
        novo = etree.SubElement(row, f"{{{NS}}}c", r=ref)
    estilo = _estilo_coluna(sheet_data, col, inicio)
    if estilo:
        novo.set("s", estilo)
    if row.get("spans"):
        del row.attrib["spans"]
    return novo


def _definir_valor(cel, valor, tipo, estilo_data=None):
    if cel.find(f"{{{NS}}}f") is not None:
        raise CelulaProtegida(f"A célula {cel.get('r')} contém fórmula e não pode ser sobrescrita.")
    for filho in list(cel):
        cel.remove(filho)
    cel.attrib.pop("t", None)
    if valor is None or valor == "":
        return
    if tipo == "data":
        if estilo_data and not cel.get("s"):
            cel.set("s", estilo_data)
        etree.SubElement(cel, f"{{{NS}}}v").text = str(_para_serial(valor))
    elif isinstance(valor, (int, float)) and not isinstance(valor, bool):
        cel.set("t", "n")
        v = int(valor) if float(valor).is_integer() else valor
        etree.SubElement(cel, f"{{{NS}}}v").text = repr(v)
    else:
        cel.set("t", "inlineStr")
        is_ = etree.SubElement(cel, f"{{{NS}}}is")
        t = etree.SubElement(is_, f"{{{NS}}}t")
        t.text = str(valor)
        if t.text != t.text.strip():
            t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")


def _atualizar_dimensao(raiz, sheet_data):
    max_lin, max_col = 1, 1
    for c in sheet_data.iter(f"{{{NS}}}c"):
        letras, lin = separa_ref(c.get("r"))
        max_lin, max_col = max(max_lin, lin), max(max_col, col_numero(letras))
    dim = raiz.find(f"{{{NS}}}dimension")
    if dim is not None:
        dim.set("ref", f"A1:{col_letra(max_col)}{max_lin}")


def _xml(raiz):
    return etree.tostring(raiz, xml_declaration=True, encoding="UTF-8", standalone=True)


def _parse(dados):
    return etree.fromstring(dados, etree.XMLParser(huge_tree=True, remove_blank_text=False))


def _estilo_celula(zin, mapa, aba, ref):
    if aba not in mapa:
        return None
    raiz = _parse(zin.read(mapa[aba]))
    for c in raiz.iter(f"{{{NS}}}c"):
        if c.get("r") == ref:
            return c.get("s")
    return None


def _verificar_desbloqueio(caminho):
    try:
        with open(caminho, "r+b"):
            pass
    except PermissionError as e:
        raise PlanilhaBloqueada(
            "A planilha está aberta no Excel (arquivo bloqueado). Feche-a e tente salvar novamente.") from e


def _reescrever(caminho, preparar, pasta_backup=None, manter_backups=40):
    """preparar(zin) -> dict {parte: bytes}; partes inexistentes são adicionadas ao pacote."""
    _verificar_desbloqueio(caminho)
    with zipfile.ZipFile(caminho) as zin:
        novos = preparar(zin)
        if not novos:
            return False
        pasta = os.path.dirname(os.path.abspath(caminho))
        fd, tmp = tempfile.mkstemp(suffix=".xlsm", dir=pasta)
        os.close(fd)
        try:
            existentes = set()
            with zipfile.ZipFile(tmp, "w") as zout:
                for info in zin.infolist():
                    existentes.add(info.filename)
                    dados = novos.get(info.filename)
                    if dados is None:
                        dados = zin.read(info.filename)
                    zout.writestr(info, dados, compress_type=info.compress_type)
                for parte, dados in novos.items():
                    if parte not in existentes:
                        zout.writestr(parte, dados, compress_type=zipfile.ZIP_DEFLATED)
        except Exception:
            os.remove(tmp)
            raise
    if pasta_backup:
        _backup(caminho, pasta_backup, manter_backups)
    try:
        os.replace(tmp, caminho)
    except PermissionError as e:
        os.remove(tmp)
        raise PlanilhaBloqueada(
            "A planilha está aberta no Excel (arquivo bloqueado). Feche-a e tente salvar novamente.") from e
    return True


def gravar_celulas(caminho, alteracoes, pasta_backup=None, manter_backups=40):
    """alteracoes: lista de (aba, ref, valor, tipo) — tipo 'data' converte ISO em data do Excel."""
    if not alteracoes:
        return 0
    por_aba = {}
    for aba, ref, valor, tipo in alteracoes:
        por_aba.setdefault(aba, []).append((ref, valor, tipo))

    def preparar(zin):
        mapa = _mapa_abas(zin)
        estilo_data = _estilo_celula(zin, mapa, ABA_CONSOLIDADO, "A6")
        novos = {}
        for aba, itens in por_aba.items():
            if aba not in mapa:
                raise KeyError(f"Aba não encontrada: {aba}")
            parte = mapa[aba]
            raiz = _parse(zin.read(parte))
            sheet_data = raiz.find(f"{{{NS}}}sheetData")
            inicio = INICIO_DADOS_ABA.get(aba, LINHA_INICIO_DADOS)
            for ref, valor, tipo in itens:
                _definir_valor(_obter_celula(sheet_data, ref, inicio), valor, tipo, estilo_data)
            _atualizar_dimensao(raiz, sheet_data)
            novos[parte] = _xml(raiz)
        wbxml = _parse(zin.read("xl/workbook.xml"))
        calc = wbxml.find(f"{{{NS}}}calcPr")
        if calc is None:
            calc = etree.SubElement(wbxml, f"{{{NS}}}calcPr")
        if calc.get("fullCalcOnLoad") != "1":
            calc.set("fullCalcOnLoad", "1")
            novos["xl/workbook.xml"] = _xml(wbxml)
        return novos

    _reescrever(caminho, preparar, pasta_backup, manter_backups)
    return len(alteracoes)


def _xml_nova_aba(spec, estilo_titulo, estilo_cab):
    def cel_txt(ref, texto, estilo):
        s = f' s="{estilo}"' if estilo else ""
        return f'<c r="{ref}"{s} t="inlineStr"><is><t>{_esc_xml(texto)}</t></is></c>'
    cols = "".join(f'<col min="{i}" max="{i}" width="{w}" customWidth="1"/>'
                   for i, (_, _, _, w) in enumerate(spec["campos"], start=1))
    cab = "".join(cel_txt(f"{col_letra(i)}2", rot, estilo_cab) for i, (_, rot, _, _) in enumerate(spec["campos"], start=1))
    ultima = col_letra(len(spec["campos"]))
    return (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        f'<worksheet xmlns="{NS}" xmlns:r="{NS_REL}">'
        f'<sheetPr><tabColor rgb="{spec["cor"]}"/></sheetPr><dimension ref="A1:{ultima}2"/>'
        '<sheetViews><sheetView workbookViewId="0"><pane ySplit="2" topLeftCell="A3" activePane="bottomLeft" state="frozen"/>'
        '<selection pane="bottomLeft" activeCell="A3" sqref="A3"/></sheetView></sheetViews>'
        f'<sheetFormatPr defaultRowHeight="15"/><cols>{cols}</cols>'
        f'<sheetData><row r="1">{cel_txt("A1", spec["titulo"], estilo_titulo)}</row>'
        f'<row r="2" ht="32" customHeight="1">{cab}</row></sheetData>'
        '<pageMargins left="0.5" right="0.5" top="0.75" bottom="0.75" header="0.3" footer="0.3"/>'
        '</worksheet>'
    ).encode("utf-8")


def _esc_xml(t):
    return str(t).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def garantir_abas(caminho, pasta_backup=None):
    """Cria as abas do sistema que ainda não existem na planilha. Devolve os nomes criados."""
    with zipfile.ZipFile(caminho) as z:
        faltando = [t for t in TABELAS.values() if t["aba"] not in _mapa_abas(z)]
    if not faltando:
        return []

    def preparar(zin):
        mapa = _mapa_abas(zin)
        estilo_titulo = _estilo_celula(zin, mapa, "CONTROLE EJ", "A1")
        estilo_cab = _estilo_celula(zin, mapa, "CONTROLE EJ", "B7")
        wbxml = _parse(zin.read("xl/workbook.xml"))
        rels = _parse(zin.read("xl/_rels/workbook.xml.rels"))
        tipos = _parse(zin.read("[Content_Types].xml"))
        sheets = wbxml.find(f"{{{NS}}}sheets")
        ids = [int(s.get("sheetId")) for s in sheets]
        rids = [r.get("Id") for r in rels]
        numeros = [int(m.group(1)) for n in zin.namelist() if (m := re.match(r"xl/worksheets/sheet(\d+)\.xml$", n))]
        prefixo = "/xl/" if any(r.get("Target", "").startswith("/") for r in rels
                                if r.get("Type", "").endswith("/worksheet")) else ""
        novos = {}
        for spec in faltando:
            n = max(numeros) + 1
            numeros.append(n)
            rid = next(f"rId{i}" for i in range(1, 10000) if f"rId{i}" not in rids)
            rids.append(rid)
            sid = max(ids) + 1
            ids.append(sid)
            parte = f"xl/worksheets/sheet{n}.xml"
            novos[parte] = _xml_nova_aba(spec, estilo_titulo, estilo_cab)
            etree.SubElement(rels, f"{{{NS_PKG_REL}}}Relationship", Id=rid,
                             Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet",
                             Target=f"{prefixo}worksheets/sheet{n}.xml")
            etree.SubElement(sheets, f"{{{NS}}}sheet", name=spec["aba"], sheetId=str(sid),
                             attrib={f"{{{NS_REL}}}id": rid})
            etree.SubElement(tipos, "{http://schemas.openxmlformats.org/package/2006/content-types}Override",
                             PartName=f"/{parte}",
                             ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml")
        novos["xl/workbook.xml"] = _xml(wbxml)
        novos["xl/_rels/workbook.xml.rels"] = _xml(rels)
        novos["[Content_Types].xml"] = _xml(tipos)
        return novos

    _reescrever(caminho, preparar, pasta_backup)
    return [t["aba"] for t in faltando]


def garantir_cabecalhos(caminho, pasta_backup=None):
    """Escreve o cabeçalho das colunas que as tabelas ganharam depois de criadas (planilhas antigas). Devolve quantos."""
    wb = openpyxl.load_workbook(caminho, read_only=True)
    try:
        faltam = []
        for spec in TABELAS.values():
            if spec["aba"] not in wb.sheetnames:
                continue
            linha2 = next(wb[spec["aba"]].iter_rows(min_row=2, max_row=2, max_col=len(spec["campos"]), values_only=True), ())
            for i, (_, rotulo, _, _) in enumerate(spec["campos"], start=1):
                if i > len(linha2) or linha2[i - 1] in (None, ""):
                    faltam.append((spec["aba"], f"{col_letra(i)}2", rotulo, None))
    finally:
        wb.close()
    if faltam:
        gravar_celulas(caminho, faltam, pasta_backup=pasta_backup)
    return len(faltam)


CATALOGO_BM = os.path.join(os.path.dirname(os.path.abspath(__file__)), "catalogo_bm.json")


def semear_catalogo_bm(caminho, pasta_backup=None):
    """Carrega o catálogo inicial de atividades do BM na aba BM ATIVIDADES, se ela estiver vazia.

    Devolve o número de atividades gravadas (0 se já havia dados ou o catálogo não existe)."""
    import json
    if not os.path.exists(CATALOGO_BM):
        return 0
    wb = openpyxl.load_workbook(caminho, read_only=True)
    try:
        if TABELAS["bm_atividades"]["aba"] not in wb.sheetnames or _ler_tabela(wb, TABELAS["bm_atividades"]):
            return 0
    finally:
        wb.close()
    with open(CATALOGO_BM, encoding="utf-8") as f:
        atividades = json.load(f)["atividades"]
    aba = TABELAS["bm_atividades"]["aba"]
    alteracoes = []
    for i, a in enumerate(atividades):
        lin = LINHA_DADOS_TABELA + i
        for col, v, t in valores_registro("bm_atividades", a):
            alteracoes.append((aba, f"{col}{lin}", v, t))
    gravar_celulas(caminho, alteracoes, pasta_backup=pasta_backup)
    return len(atividades)


BM_MEDIDO = os.path.join(os.path.dirname(os.path.abspath(__file__)), "bm_medido_bm1.json")


def registrar_bm_medido(caminho, dados, empresas=None, refazer=False, com_bm2_rascunho=False, pasta_backup=None):
    """Registra como fechado (medido e pago) o BM1 de cada empresa: corte, quantidade acumulada medida por atividade e
    descontos do BM1. `dados` vem de bm_medido_bm1.json. Devolve o número de linhas gravadas por tabela."""
    empresas = [e for e in (empresas or dados["empresas"]) if e in dados["empresas"]]
    T = ler_planilha(caminho)["tabelas"]
    catalogo = {}
    for a in T["bm_atividades"]:
        catalogo.setdefault(a["empresa"], []).append(a)
    celulas = []

    def limpar(tabela, linhas):
        spec = TABELAS[tabela]
        for r in linhas:
            celulas.extend((spec["aba"], f"{col_letra(i)}{r['linha']}", None, None) for i in range(1, len(spec["campos"]) + 1))

    antigos = {t: [r for r in T[t] if r["empresa"] in empresas and int(r.get("bm") or 0) == 1]
               for t in ("bm_periodos", "bm_fechamentos", "bm_fech_atividades")}
    ant_ded = [r for r in T["bm_deducoes"] if r["empresa"] in empresas and (r.get("descricao") or "").startswith("[planilha]")]
    if (any(antigos.values()) or ant_ded) and not refazer:
        raise ValueError("O BM1 já está registrado na planilha.")
    for t, linhas in antigos.items():
        limpar(t, linhas)
    limpar("bm_deducoes", ant_ded)
    reservadas = {**antigos, "bm_deducoes": ant_ded}

    novos = {"bm_periodos": [], "bm_fechamentos": [], "bm_fech_atividades": [], "bm_deducoes": []}
    for emp in empresas:
        d = dados["empresas"][emp]
        novos["bm_periodos"].append({"empresa": emp, "bm": 1, "corte": dados["corte"], "obs": "BM1 pago (planilha de medição)"})
        novos["bm_fechamentos"].append({"empresa": emp, "bm": 1, "fechado_em": dados["fechado_em"], "equip": 0, "comb": 0, "litros": 0,
                                        "obs": f"BM1 já medido e pago: R$ {d['medido']:,.2f} medido, R$ {d['a_faturar']:,.2f} faturado (descontos em BM DEDUÇÕES)"})
        for a in catalogo.get(emp, []):
            novos["bm_fech_atividades"].append({"empresa": emp, "bm": 1, "codigo": a["codigo"], "item": a["item"],
                                                "realizado": d["realizado"].get(a["codigo"], 0.0), "qtd": a["qtd"], "peso": a["peso"],
                                                "valor_qpc": a.get("valor_qpc"), "valor_rotula": a.get("valor_rotula")})
        for l in d["linhas"]:
            novos["bm_deducoes"].append({"empresa": emp, "bm": 1, "data": dados["corte"], "tipo": l["tipo"], "valor": l["valor"],
                                         "percentual": l["percentual"], "contrato": "RÓTULA", "descricao": "[planilha] " + l["descricao"]})
    if com_bm2_rascunho:
        for emp, d in (dados.get("bm2_rascunho") or {}).items():
            if emp not in empresas:
                continue
            for l in d["linhas"]:
                novos["bm_deducoes"].append({"empresa": emp, "bm": 2, "data": dt.date.today().isoformat(), "tipo": l["tipo"], "valor": l["valor"],
                                             "percentual": l["percentual"], "contrato": "RÓTULA",
                                             "descricao": "[planilha] " + l["descricao"] + " (rascunho do BM2: conferir)"})
    # cabeçalho da coluna "% do medido" (planilhas criadas antes do campo)
    spec_d = TABELAS["bm_deducoes"]
    celulas.append((spec_d["aba"], f"{col_letra(len(spec_d['campos']))}2", spec_d["campos"][-1][1], None))
    for tabela, regs in novos.items():
        spec = TABELAS[tabela]
        livres_usadas = {r["linha"] for r in T[tabela]} - {r["linha"] for r in reservadas.get(tabela, [])}
        lin = LINHA_DADOS_TABELA
        for campos in regs:
            while lin in livres_usadas:
                lin += 1
            for col, v, tp in valores_registro(tabela, campos):
                celulas.append((spec["aba"], f"{col}{lin}", v, tp))
            lin += 1
    gravar_celulas(caminho, celulas, pasta_backup=pasta_backup)
    return {t: len(r) for t, r in novos.items()}


def semear_bm_medido(caminho, pasta_backup=None):
    """Na primeira abertura, registra o BM1 já pago das empresas que ainda não têm nenhum BM cadastrado
    (nem corte nem fechamento). Devolve as empresas registradas."""
    import json
    if not os.path.exists(BM_MEDIDO):
        return []
    with open(BM_MEDIDO, encoding="utf-8") as f:
        dados = json.load(f)
    T = ler_planilha(caminho)["tabelas"]
    if not T["bm_atividades"]:
        return []
    com_bm = {r["empresa"] for r in T["bm_periodos"]} | {r["empresa"] for r in T["bm_fechamentos"]}
    novas = [e for e in dados["empresas"] if e not in com_bm and any(a["empresa"] == e for a in T["bm_atividades"])]
    if not novas:
        return []
    registrar_bm_medido(caminho, dados, empresas=novas, pasta_backup=pasta_backup)
    return novas


def _backup(caminho, pasta, manter):
    os.makedirs(pasta, exist_ok=True)
    base, ext = os.path.splitext(os.path.basename(caminho))
    carimbo = dt.datetime.now().strftime("%Y%m%d_%H%M%S")
    shutil.copy2(caminho, os.path.join(pasta, f"{base}_{carimbo}{ext}"))
    antigos = sorted(f for f in os.listdir(pasta) if f.startswith(base + "_"))
    for f in antigos[:-manter]:
        os.remove(os.path.join(pasta, f))


# ------------------------------------------- importação das abas mensais de equipamento

def _normalizar_empresa(v):
    t = re.sub(r"\s+", " ", str(v or "")).strip().upper()
    if not t:
        return None
    partes = []
    for p in re.split(r"\s*/\s*", t):
        if p in ("GA",) or p.startswith("GLOBO"):
            p = "GLOBO AÇOS"
        partes.append(p)
    return "/".join(partes)


def _avaliar_formula(f, linha_vals):
    """Avalia fórmulas aritméticas simples (ex.: =7.58*O5, =3000+522) usando valores da mesma linha."""
    import ast
    import operator as op
    expr = f.lstrip("=")

    def troca(m):
        v = linha_vals.get(m.group(1))
        n = numero(v)
        return repr(n if n is not None else 0)
    expr = re.sub(r"\$?([A-Z]{1,3})\$?\d+", troca, expr)
    ops = {ast.Add: op.add, ast.Sub: op.sub, ast.Mult: op.mul, ast.Div: op.truediv, ast.USub: op.neg}

    def ev(n):
        if isinstance(n, ast.Expression):
            return ev(n.body)
        if isinstance(n, ast.Constant) and isinstance(n.value, (int, float)):
            return n.value
        if isinstance(n, ast.BinOp) and type(n.op) in ops:
            return ops[type(n.op)](ev(n.left), ev(n.right))
        if isinstance(n, ast.UnaryOp) and type(n.op) in ops:
            return ops[type(n.op)](ev(n.operand))
        raise ValueError(f)
    try:
        return ev(ast.parse(expr, mode="eval"))
    except (ValueError, SyntaxError, ZeroDivisionError):
        return None


def _preco_da_formula(f):
    m = re.match(r"=\s*([\d.]+)\s*\*", str(f or ""))
    return float(m.group(1)) if m else None


def _nome_equip(titulo):
    t = re.sub(r"\s+", " ", str(titulo or "")).strip().upper()
    if "SANY" in t:
        return "SANY"
    return t or None


def extrair_usos_mensais(caminho):
    """Lê as abas 'USO DE EQUIPAMENTO <MÊS>' (blocos lado a lado: título na linha 3, cabeçalho na 4)."""
    wb = openpyxl.load_workbook(caminho, read_only=True)
    registros = []
    try:
        for aba in [n for n in wb.sheetnames if n.upper().startswith("USO DE EQUIPAMENTO")]:
            linhas = _linhas(wb[aba])
            if len(linhas) < 5:
                continue
            cab = {c: str(_cel(linhas, 4, c) or "").strip().upper() for c in range(1, len(linhas[3]) + 1)}
            inicios = [c for c, h in cab.items() if h == "DATA"]
            for i, c0 in enumerate(inicios):
                c1 = (inicios[i + 1] - 1) if i + 1 < len(inicios) else max(cab)
                titulo = next((_cel(linhas, 3, c) for c in range(c0, c1 + 1) if _cel(linhas, 3, c)), None)
                if not titulo:
                    continue
                cols = {cab[c]: c for c in range(c0 + 1, c1 + 1) if cab[c]}
                abastecimento = "ABASTEC" in str(titulo).upper()
                col_emp = cols.get("EMPRESA") or cols.get("USO")
                col_val = cols.get("VALOR")
                col_obs = next((c for c in range(c0 + 1, c1 + 1) if not cab[c]), None)
                if abastecimento and not cols.get("LITROS"):
                    continue
                for lin in range(5, len(linhas) + 1):
                    data = iso(_cel(linhas, lin, c0))
                    if not data:
                        continue
                    vals = {col_letra(c): _cel(linhas, lin, c) for c in range(1, len(linhas[lin - 1]) + 1)}
                    empresa = _normalizar_empresa(_cel(linhas, lin, col_emp)) if col_emp else None
                    bruto = _cel(linhas, lin, col_val) if col_val else None
                    valor = _avaliar_formula(bruto, vals) if eh_formula(bruto) else numero(bruto)
                    reg = {"data": data, "empresa": empresa, "origem": f"{aba}!{col_letra(c0)}{lin}",
                           "operador": _cel(linhas, lin, cols["OPERADOR"]) if "OPERADOR" in cols else None,
                           "obs": _cel(linhas, lin, col_obs) if col_obs and not abastecimento else None}
                    if abastecimento:
                        litros = numero(_cel(linhas, lin, cols["LITROS"]))
                        if not litros:
                            continue
                        preco = _preco_da_formula(bruto) if eh_formula(bruto) else (valor / litros if valor else None)
                        reg.update(equipamento=_nome_equip(_cel(linhas, lin, cols.get("MAQUINA", 0)) if cols.get("MAQUINA") else None),
                                   uso="ABASTECIMENTO", litros=litros, preco_litro=preco,
                                   custo_combustivel=round(valor, 2) if valor else (round(litros * preco, 2) if preco else None))
                    else:
                        if not valor or not empresa or empresa in ("FERIADO", "NÃO USO", "CHEGADA"):
                            continue
                        reg.update(equipamento=_nome_equip(titulo), uso="DIÁRIA", quantidade=1, custo=round(valor, 2))
                    registros.append(reg)
    finally:
        wb.close()
    return registros


# ------------------------------------------- planilha de estoque (somente leitura)

ABAS_ESTOQUE = ("Consumo_Joists", "Controle_Materiais")


def eh_planilha_producao(caminho):
    try:
        wb = openpyxl.load_workbook(caminho, read_only=True)
        try:
            return ABA_CONSOLIDADO in wb.sheetnames
        finally:
            wb.close()
    except Exception:
        return False


def eh_planilha_estoque(caminho):
    try:
        wb = openpyxl.load_workbook(caminho, read_only=True)
        try:
            return all(a in wb.sheetnames for a in ABAS_ESTOQUE)
        finally:
            wb.close()
    except Exception:
        return False


def _txt(v):
    if v is None:
        return None
    t = str(v).strip()
    return None if t == "" or t.startswith("#") else t


def ler_estoque_externo(caminho):
    """Lê os indicadores já calculados pela planilha de estoque (valores salvos pelo Excel)."""
    wb = openpyxl.load_workbook(caminho, read_only=True, data_only=True)
    try:
        cj = _linhas(wb["Consumo_Joists"])
        mix = []
        if "Apontamento_Montagem" in wb.sheetnames:
            am = _linhas(wb["Apontamento_Montagem"], max_col=10)
            for lin in range(4, 14):
                tipo = _txt(_cel(am, lin, 1))
                if tipo:
                    mix.append({"tipo": tipo, "planejada": numero(_cel(am, lin, 3)) or 0,
                                "premontadas": numero(_cel(am, lin, 4)) or 0, "montadas": numero(_cel(am, lin, 5)) or 0,
                                "pendentes": numero(_cel(am, lin, 6)) or 0, "aguardando_icamento": numero(_cel(am, lin, 8)) or 0})
        pesos = [m["pendentes"] for m in mix][:10] or [1]
        if not any(pesos):
            pesos = [1] + [0] * 9

        itens, secao = [], "PREMONTAGEM"
        for lin in range(4, len(cj) + 1):
            tag = _txt(_cel(cj, lin, 1))
            if not tag:
                continue
            if tag.upper().startswith("CONSUMO DE FIX"):
                continue
            if tag.upper() == "TAG":
                continue
            por_tipo = [numero(_cel(cj, lin, c)) or 0 for c in range(5, 15)]
            perda = numero(_cel(cj, lin, 15)) or 0
            soma_p = sum(pesos[:len(por_tipo)]) or 1
            coef = sum(p * q for p, q in zip(pesos, por_tipo)) / soma_p * (1 + perda)
            etapa = "IÇAMENTO JOIST" if tag.upper().startswith("FG") else "PREMONTAGEM"
            itens.append({
                "tag": tag, "produto": _txt(_cel(cj, lin, 2)), "descricao": _txt(_cel(cj, lin, 3)),
                "material": _txt(_cel(cj, lin, 4)), "etapa": etapa, "coef": round(coef, 4), "perda": perda,
                "recebido": numero(_cel(cj, lin, 18)), "consumido": numero(_cel(cj, lin, 19)),
                "disponivel": numero(_cel(cj, lin, 20)), "minimo": numero(_cel(cj, lin, 21)),
                "status": _txt(_cel(cj, lin, 27)), "solicitar": numero(_cel(cj, lin, 28)),
            })

        capacidade = {}
        for etapa, col in (("PREMONTAGEM", 37), ("IÇAMENTO JOIST", 39)):
            capacidade[etapa] = {"joists": numero(_cel(cj, 4, col)), "limitante": _txt(_cel(cj, 5, col)),
                                 "descricao": _txt(_cel(cj, 6, col))}

        fixadores = []
        if "Consumo_Quadrantes" in wb.sheetnames:
            cq = _linhas(wb["Consumo_Quadrantes"], max_col=12)
            for lin in range(5, len(cq) + 1):
                cod = _txt(_cel(cq, lin, 1))
                if not cod or not re.match(r"^[A-Z]{2}\d", cod):
                    if fixadores:
                        break
                    continue
                fixadores.append({
                    "codigo": cod, "descricao": _txt(_cel(cq, lin, 2)), "escopo": _txt(_cel(cq, lin, 3)),
                    "consumo_total": numero(_cel(cq, lin, 6)), "disponivel": numero(_cel(cq, lin, 7)),
                    "saldo": numero(_cel(cq, lin, 8)), "status": _txt(_cel(cq, lin, 10)), "regra": _txt(_cel(cq, lin, 11)),
                })

        inventario = []
        if "INVENTARIO" in wb.sheetnames:
            inv = _linhas(wb["INVENTARIO"], max_col=16)
            for lin in range(2, len(inv) + 1):
                tag = _txt(_cel(inv, lin, 3))
                virtual, fisico = numero(_cel(inv, lin, 9)), numero(_cel(inv, lin, 6))
                if not tag or not (virtual or fisico):
                    continue
                inventario.append({"tag": tag, "produto": _txt(_cel(inv, lin, 2)), "virtual": virtual, "fisico": fisico,
                                   "perda": numero(_cel(inv, lin, 14)), "acuracidade": numero(_cel(inv, lin, 16))})
    finally:
        wb.close()
    return {
        "arquivo": os.path.basename(caminho),
        "modificado": dt.datetime.fromtimestamp(os.path.getmtime(caminho)).isoformat(timespec="seconds"),
        "itens": itens, "capacidade": capacidade, "mix_joists": mix,
        "fixadores": fixadores, "inventario": inventario,
    }
