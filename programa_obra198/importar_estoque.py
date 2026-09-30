"""Importa dados de estoque das abas Controle_Materiais, REMESSAS, CONSUMO FISICO e INVENTARIO
da planilha teste claude.xlsm para dados/estoque_obra198.json.

Uso:  python importar_estoque.py [caminho_planilha.xlsm]
"""
import json
import os
import sys

import openpyxl

BASE = os.path.dirname(os.path.abspath(__file__))
RAIZ = os.path.dirname(BASE)
SAIDA = os.path.join(RAIZ, "dados", "estoque_obra198.json")


def _v(ws, r, c):
    v = ws.cell(r, c).value
    return v


def _s(v):
    return str(v).strip() if v not in (None, "") else ""


def _n(v):
    if v is None or v == "":
        return 0
    try:
        return float(v)
    except (ValueError, TypeError):
        return 0


def importar(caminho):
    print(f"Lendo {caminho}...")
    wb = openpyxl.load_workbook(caminho, read_only=True, data_only=True)

    # === 1. CONTROLE MATERIAIS ===
    ws = wb["Controle_Materiais"]
    materiais = []
    for r in range(3, ws.max_row + 1):
        tag = _v(ws, r, 1)
        if not tag:
            continue
        materiais.append({
            "tag": _s(tag),
            "peso_unitario": _n(_v(ws, r, 2)),
            "peso_total": _n(_v(ws, r, 3)),
            "comp_mm": _n(_v(ws, r, 4)),
            "material": _s(_v(ws, r, 5)),
            "produto": _s(_v(ws, r, 6)),
            "planejado": _n(_v(ws, r, 8)),
            "chegou": _n(_v(ws, r, 10)),
            "necessario_real": _n(_v(ws, r, 11)),
            "falta_enviar": _n(_v(ws, r, 13)),
            "consumido": _n(_v(ws, r, 15)),
            "estoque_pos_baixa": _n(_v(ws, r, 16)),
            "atendimento": _n(_v(ws, r, 19)),
            "status_logistico": _s(_v(ws, r, 20)),
            "tipo_material": _s(_v(ws, r, 21)),
            "prioridade": _s(_v(ws, r, 22)),
            "etapa": _s(_v(ws, r, 23)),
        })
    print(f"  Controle_Materiais: {len(materiais)} itens")

    # === 2. REMESSAS ===
    ws2 = wb["REMESSAS"]
    rem_nomes = []
    for c in range(5, ws2.max_column + 1):
        v = _v(ws2, 2, c)
        if not v:
            continue
        s = _s(v)
        if "TOTAL" in s.upper() or "PLANEJADO" in s.upper() or "FALTA" in s.upper():
            break
        rem_nomes.append({"col": c, "nome": s})

    remessas_itens = []
    for r in range(3, ws2.max_row + 1):
        tag = _v(ws2, r, 3)
        if not tag:
            continue
        row = {
            "material": _s(_v(ws2, r, 1)),
            "produto": _s(_v(ws2, r, 2)),
            "tag": _s(tag),
            "etapa": _s(_v(ws2, r, 4)),
            "qtd_por_remessa": {},
            "total_recebido": 0,
        }
        for rc in rem_nomes:
            v = _v(ws2, r, rc["col"])
            q = _n(v)
            if q > 0:
                row["qtd_por_remessa"][rc["nome"]] = q
                row["total_recebido"] += q
        remessas_itens.append(row)
    print(f"  REMESSAS: {len(remessas_itens)} itens, {len(rem_nomes)} remessas")

    # === 3. CONSUMO FISICO ===
    ws3 = wb["CONSUMO FISICO"]
    sem_cols = []
    c = 3
    while c <= 26:
        s = _v(ws3, 1, c)
        if s:
            txt = _s(s)
            if "TOTAL" not in txt.upper() and "ESTOQUE" not in txt.upper():
                sem_cols.append({"col_ej": c, "col_cmm": c + 1, "semana": txt})
        c += 2

    inv_cols = []
    for c in range(27, ws3.max_column + 1):
        h1 = _v(ws3, 1, c)
        h2 = _v(ws3, 2, c)
        if h1:
            inv_cols.append({"col": c, "label": _s(h1)})

    consumo_itens = []
    for r in range(3, ws3.max_row + 1):
        tag = _v(ws3, r, 2)
        if not tag:
            continue
        row = {
            "produto": _s(_v(ws3, r, 1)),
            "tag": _s(tag),
            "consumo_semanas": {},
            "total_consumo": _n(_v(ws3, r, 23)),
            "inventarios_fisicos": {},
        }
        for sc in sem_cols:
            ej = _n(_v(ws3, r, sc["col_ej"]))
            cmm = _n(_v(ws3, r, sc["col_cmm"]))
            if ej or cmm:
                row["consumo_semanas"][sc["semana"]] = {"EJ": ej, "CMM": cmm, "total": ej + cmm}
        for ic in inv_cols:
            v = _v(ws3, r, ic["col"])
            if v is not None:
                row["inventarios_fisicos"][ic["label"]] = _n(v)
        consumo_itens.append(row)
    print(f"  CONSUMO FISICO: {len(consumo_itens)} itens, {len(sem_cols)} semanas")

    # === 4. INVENTARIO ===
    ws4 = wb["INVENTARIO"]
    inventario_itens = []
    for r in range(2, ws4.max_row + 1):
        tag = _v(ws4, r, 3)
        if not tag:
            continue
        inventario_itens.append({
            "material": _s(_v(ws4, r, 1)),
            "produto": _s(_v(ws4, r, 2)),
            "tag": _s(tag),
            "planejado": _n(_v(ws4, r, 4)),
            "estoque_virtual_base": _n(_v(ws4, r, 5)),
            "estoque_fisico": _n(_v(ws4, r, 6)),
            "consumo_virtual": _n(_v(ws4, r, 7)),
            "consumo_fisico": _n(_v(ws4, r, 8)),
            "estoque_virtual_atual": _n(_v(ws4, r, 9)),
            "dif_estoque": _n(_v(ws4, r, 10)),
            "perda_real": _n(_v(ws4, r, 14)),
            "taxa_perda": _n(_v(ws4, r, 15)),
            "acuracidade": _n(_v(ws4, r, 16)),
            "status": _s(_v(ws4, r, 17)),
            "acao": _s(_v(ws4, r, 18)),
            "obs": _s(_v(ws4, r, 19)),
        })
    print(f"  INVENTARIO: {len(inventario_itens)} itens")

    wb.close()

    resultado = {
        "fonte": os.path.basename(caminho),
        "materiais": materiais,
        "remessas": {
            "colunas": [rc["nome"] for rc in rem_nomes],
            "itens": remessas_itens,
        },
        "consumo_fisico": {
            "semanas": [sc["semana"] for sc in sem_cols],
            "empresas": ["EJ", "CMM"],
            "itens": consumo_itens,
            "inventarios_datas": [ic["label"] for ic in inv_cols],
        },
        "inventario": inventario_itens,
    }

    os.makedirs(os.path.dirname(SAIDA), exist_ok=True)
    with open(SAIDA, "w", encoding="utf-8") as f:
        json.dump(resultado, f, ensure_ascii=False, indent=2, default=str)
    print(f"Exportado para {SAIDA}")
    return resultado


if __name__ == "__main__":
    planilha = sys.argv[1] if len(sys.argv) > 1 else os.path.join(RAIZ, "teste claude.xlsm")
    if not os.path.exists(planilha):
        sys.exit(f"Planilha não encontrada: {planilha}")
    importar(planilha)
