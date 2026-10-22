import os
import sys
import tempfile
import unittest
import zipfile

import openpyxl

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
import validadores  # noqa: E402

RAIZ = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))


def pdf_com_texto(texto):
    """PDF mínimo de 1 página com texto selecionável."""
    conteudo = "BT /F1 12 Tf 50 700 Td (%s) Tj ET" % texto.replace("(", "").replace(")", "")
    objs = ["<< /Type /Catalog /Pages 2 0 R >>", "<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
            "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
            "<< /Length %d >>\nstream\n%s\nendstream" % (len(conteudo), conteudo), "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"]
    out, pos = "%PDF-1.4\n", []
    for i, o in enumerate(objs, 1):
        pos.append(len(out.encode("latin-1")))
        out += "%d 0 obj\n%s\nendobj\n" % (i, o)
    xref = len(out.encode("latin-1"))
    out += "xref\n0 %d\n0000000000 65535 f \n" % (len(objs) + 1) + "".join("%010d 00000 n \n" % p for p in pos)
    out += "trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objs) + 1, xref)
    return out.encode("latin-1")


class Base(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.mkdtemp()

    def arq(self, nome, dados=b""):
        p = os.path.join(self.tmp, nome)
        with open(p, "wb") as f:
            f.write(dados)
        return p


class TesteIfc(Base):
    def test_ifc_real(self):
        p = os.path.join(RAIZ, "LSF-MET-EX-230-200-BIM-R0A.ifc")
        if not os.path.exists(p):
            self.skipTest("IFC de exemplo ausente")
        r = validadores.validar("ifc", p, "x.ifc")
        self.assertTrue(r["ok"], r)
        self.assertGreater(r["resumo"]["conjuntos"], 100)
        self.assertGreater(r["resumo"]["codigos_de_posicao"], 50)
        self.assertGreater(r["resumo"]["eixos_da_malha"], 0)
        self.assertTrue(r["resumo"]["sistema_tekla"])

    def test_ifc_dentro_de_zip(self):
        p = os.path.join(RAIZ, "LSF-MET-EX-235-200-BIM-R0A.ifc")
        if not os.path.exists(p):
            self.skipTest("IFC de exemplo ausente")
        z = os.path.join(self.tmp, "m.zip")
        with zipfile.ZipFile(z, "w") as zz:
            zz.write(p, "modelo.ifc")
        r = validadores.validar("ifc", z, "m.zip")
        self.assertTrue(r["ok"], r)
        self.assertEqual(r["resumo"]["arquivo"], "modelo.ifc")

    def test_nao_e_ifc(self):
        r = validadores.validar("ifc", self.arq("a.ifc", b"nada a ver"), "a.ifc")
        self.assertFalse(r["ok"])

    def test_zip_sem_ifc(self):
        z = os.path.join(self.tmp, "v.zip")
        with zipfile.ZipFile(z, "w") as zz:
            zz.writestr("leia.txt", "oi")
        self.assertFalse(validadores.validar("ifc", z, "v.zip")["ok"])


class TesteFixadores(Base):
    def test_xlsx_ok(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["Relatório de parafusos"])
        ws.append(["Descrição", "Diâmetro", "Comprimento", "Norma", "Quantidade"])
        ws.append(["PARAF. 1/2 x 1 1/4", '1/2"', 31.75, "A325", 120])
        ws.append(["PARAF. 5/8 x 2", '5/8"', 50, "A307", 80])
        p = os.path.join(self.tmp, "f.xlsx")
        wb.save(p)
        r = validadores.validar("fixadores", p, "f.xlsx")
        self.assertTrue(r["ok"], r)
        self.assertEqual(r["resumo"]["total_de_pecas"], 200)

    def test_xlsx_sem_quantidade(self):
        wb = openpyxl.Workbook()
        wb.active.append(["a", "b"])
        p = os.path.join(self.tmp, "f.xlsx")
        wb.save(p)
        self.assertFalse(validadores.validar("fixadores", p, "f.xlsx")["ok"])

    def test_csv_ok(self):
        r = validadores.validar("fixadores", self.arq("f.csv", "Descrição;Diâmetro;Quantidade\nPARAF;1/2;10\n".encode()), "f.csv")
        self.assertTrue(r["ok"], r)

    def test_trb_aceita_com_aviso(self):
        r = validadores.validar("fixadores", self.arq("m.trb", b"x" * 5000), "m.trb")
        self.assertTrue(r["ok"])
        self.assertTrue(r["avisos"])


class TesteCronograma(Base):
    XML = ('<?xml version="1.0"?><Project xmlns="http://schemas.microsoft.com/project"><Tasks>'
           '<Task><Name>Fundação</Name><Start>2026-03-31T08:00:00</Start><Finish>2026-05-10T17:00:00</Finish></Task>'
           '<Task><Name>Montagem</Name><Start>2026-05-11T08:00:00</Start><Finish>2026-09-01T17:00:00</Finish></Task></Tasks></Project>')

    def test_xml_ok(self):
        r = validadores.validar("cronograma", self.arq("c.xml", self.XML.encode()), "c.xml")
        self.assertTrue(r["ok"], r)
        self.assertEqual(r["resumo"]["tarefas"], 2)
        self.assertEqual(r["resumo"]["fim"], "2026-09-01")

    def test_xml_errado(self):
        self.assertFalse(validadores.validar("cronograma", self.arq("c.xml", b"<a><b/></a>"), "c.xml")["ok"])

    def test_xlsx_ok(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["Nome da tarefa", "Início", "Término"])
        ws.append(["Montagem de vigas", "2026-05-01", "2026-06-01"])
        p = os.path.join(self.tmp, "c.xlsx")
        wb.save(p)
        self.assertTrue(validadores.validar("cronograma", p, "c.xlsx")["ok"])

    def test_pdf_com_datas(self):
        txt = " ".join("Tarefa %d 0%d/0%d/26 1%d/0%d/26" % (i, i % 9 + 1, i % 9 + 1, i % 9, i % 9 + 1) for i in range(8))
        r = validadores.validar("cronograma", self.arq("c.pdf", pdf_com_texto(txt)), "c.pdf")
        self.assertTrue(r["ok"], r)
        self.assertTrue(r["avisos"])

    def test_pdf_sem_texto(self):
        r = validadores.validar("cronograma", self.arq("c.pdf", pdf_com_texto("")), "c.pdf")
        self.assertFalse(r["ok"])


class TesteOrcamento(Base):
    def test_ok(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "BM - EJ - RÓTULA"
        ws.append(["ÍTEM", "DESCRIÇÃO", "QUANTIDADE", "UNIDADE", "VALOR UNITÁRIO", "VALOR"])
        ws.append(["3.1.1", "Instalação de joists", 817, "un", 100.0, 81700.0])
        ws.append(["3.1.2", "Vigas", 120, "un", 50.0, 6000.0])
        p = os.path.join(self.tmp, "o.xlsx")
        wb.save(p)
        r = validadores.validar("orcamento", p, "o.xlsx")
        self.assertTrue(r["ok"], r)
        self.assertEqual(r["resumo"]["itens"], 2)
        self.assertEqual(r["resumo"]["valor_somado"], 87700.0)

    def test_sem_tabela(self):
        wb = openpyxl.Workbook()
        wb.active.append(["foo", "bar"])
        p = os.path.join(self.tmp, "o.xlsx")
        wb.save(p)
        self.assertFalse(validadores.validar("orcamento", p, "o.xlsx")["ok"])


class TesteProjetos(Base):
    def test_pdf_ok(self):
        r = validadores.validar("projetos", self.arq("p.pdf", pdf_com_texto("Cobertura telha trapezoidal 6000 m2 vao 12 m")), "p.pdf")
        self.assertTrue(r["ok"], r)
        self.assertGreater(r["resumo"]["cita_telha"], 0)

    def test_pdf_sem_texto(self):
        self.assertFalse(validadores.validar("projetos", self.arq("p.pdf", pdf_com_texto("")), "p.pdf")["ok"])

    def test_nao_pdf(self):
        self.assertFalse(validadores.validar("projetos", self.arq("p.pdf", b"nao sou pdf"), "p.pdf")["ok"])


if __name__ == "__main__":
    unittest.main()
