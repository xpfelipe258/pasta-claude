import openpyxl
from openpyxl.utils import get_column_letter
from copy import copy

SRC = "/root/.claude/uploads/2698103b-beaa-5416-9ed3-fc54bf3f5799/56216af6-PLANO_DE_PRODU__O_198_POR_EMPRESA_V2.xlsm"
DST = "/home/user/pasta-claude/PLANO_DE_PRODUCAO_198_POR_EMPRESA_V3.xlsm"

wb = openpyxl.load_workbook(SRC, keep_vba=True)

# ─────────────────────────────────────────────────────────
# 1. BASE METAS DIÁRIAS — add GLOBO AÇOS columns (U-AC)
# ─────────────────────────────────────────────────────────
ws_bmd = wb['BASE METAS DIÁRIAS']

headers = {
    21: 'GA PERF', 22: 'GA TELHAR', 23: 'GA FECH',
    24: 'DU GA PERF', 25: 'DU GA TELHAR', 26: 'DU GA FECH',
    27: 'SEM GA PERF', 28: 'SEM GA TELHAR', 29: 'SEM GA FECH',
}
for col, hdr in headers.items():
    ws_bmd.cell(row=1, column=col).value = hdr

GA_META = [19, 7, 0]
GA_DU = 5

for row_idx in range(2, ws_bmd.max_row + 1):
    ws_bmd.cell(row=row_idx, column=21).value = GA_META[0]
    ws_bmd.cell(row=row_idx, column=22).value = GA_META[1]
    ws_bmd.cell(row=row_idx, column=23).value = GA_META[2]
    ws_bmd.cell(row=row_idx, column=24).value = GA_DU
    ws_bmd.cell(row=row_idx, column=25).value = GA_DU
    ws_bmd.cell(row=row_idx, column=26).value = GA_DU
    ws_bmd.cell(row=row_idx, column=27).value = GA_META[0] * GA_DU  # 95
    ws_bmd.cell(row=row_idx, column=28).value = GA_META[1] * GA_DU  # 35
    ws_bmd.cell(row=row_idx, column=29).value = GA_META[2] * GA_DU  # 0

print("1/5 BASE METAS DIÁRIAS: +9 colunas (U-AC) GLOBO AÇOS")

# ─────────────────────────────────────────────────────────
# 2. PARAMETROS KPI — add GLOBO AÇOS rows 11-13
# ─────────────────────────────────────────────────────────
ws_param = wb['PARAMETROS KPI']

ga_items = [
    ('GLOBO AÇOS', 'PERFILAÇÃO'),
    ('GLOBO AÇOS', 'TELHAR'),
    ('GLOBO AÇOS', 'FECHAMENTO'),
]

for i, (empresa, servico) in enumerate(ga_items):
    r = 11 + i
    ws_param.cell(row=r, column=1).value = empresa
    ws_param.cell(row=r, column=2).value = servico
    for col_idx in range(3, 6):
        col_letter = get_column_letter(col_idx)
        ws_param.cell(row=r, column=col_idx).value = (
            f"=SUMIFS('METAS EMPRESAS'!${col_letter}$5:${col_letter}$13,"
            f"'METAS EMPRESAS'!$A$5:$A$13,A{r},"
            f"'METAS EMPRESAS'!$B$5:$B$13,B{r})"
        )
    ws_param.cell(row=r, column=6).value = 'METAS EMPRESAS'
    ws_param.cell(row=r, column=7).value = 'ATIVA'

# Also update existing EJ/CMM rows (5-10) to reference $5:$13 range
for r in range(5, 11):
    for col_idx in range(3, 6):
        col_letter = get_column_letter(col_idx)
        ws_param.cell(row=r, column=col_idx).value = (
            f"=SUMIFS('METAS EMPRESAS'!${col_letter}$5:${col_letter}$13,"
            f"'METAS EMPRESAS'!$A$5:$A$13,A{r},"
            f"'METAS EMPRESAS'!$B$5:$B$13,B{r})"
        )

print("2/5 PARAMETROS KPI: +3 linhas (11-13) GLOBO AÇOS, ranges $5:$13")

# ─────────────────────────────────────────────────────────
# 3. CONTROLE KPI — append GLOBO AÇOS (1071 rows)
# ─────────────────────────────────────────────────────────
ws_kpi = wb['CONTROLE KPI']
OLD_MAX = 2143
NEW_START = 2144

dates = []
for row_idx in range(2, OLD_MAX + 1, 6):
    d = ws_kpi.cell(row=row_idx, column=1).value
    if d is not None:
        dates.append(d)

ga_services = ['PERFILAÇÃO', 'TELHAR', 'FECHAMENTO']

for date_idx, date_val in enumerate(dates):
    for svc_idx, servico in enumerate(ga_services):
        r = NEW_START + date_idx * 3 + svc_idx
        f_col = svc_idx + 1

        ws_kpi.cell(row=r, column=1).value = date_val
        ws_kpi.cell(row=r, column=2).value = 'GERAL'
        ws_kpi.cell(row=r, column=3).value = 'GLOBO AÇOS'
        ws_kpi.cell(row=r, column=4).value = servico

        # F: Meta semanal (AA-AC)
        ws_kpi.cell(row=r, column=6).value = (
            f"=IFERROR(INDEX('BASE METAS DIÁRIAS'!$AA$2:$AC$358,"
            f"MATCH($A{r},'BASE METAS DIÁRIAS'!$A$2:$A$358,0),{f_col}),\"\")"
        )
        # G: Previsto diário (U-W / X-Z)
        ws_kpi.cell(row=r, column=7).value = (
            f"=IFERROR(IF(WEEKDAY($A{r},2)<="
            f"INDEX('BASE METAS DIÁRIAS'!$X$2:$Z$358,"
            f"MATCH($A{r},'BASE METAS DIÁRIAS'!$A$2:$A$358,0),{f_col}),"
            f"INDEX('BASE METAS DIÁRIAS'!$U$2:$W$358,"
            f"MATCH($A{r},'BASE METAS DIÁRIAS'!$A$2:$A$358,0),{f_col}),\"\"),\"\")"
        )
        # H: Realizado (consolidado H:J)
        ws_kpi.cell(row=r, column=8).value = (
            f"=IFERROR(IF(INDEX('CONTROLE PRODUÇÃO POR EMPRESAS'!$H$6:$J$5000,"
            f"MATCH($A{r},'CONTROLE PRODUÇÃO POR EMPRESAS'!$A$6:$A$5000,0),"
            f"{f_col})=\"\",\"\","
            f"IFERROR(1*INDEX('CONTROLE PRODUÇÃO POR EMPRESAS'!$H$6:$J$5000,"
            f"MATCH($A{r},'CONTROLE PRODUÇÃO POR EMPRESAS'!$A$6:$A$5000,0),"
            f"{f_col}),\"\")),\"\")"
        )
        # I: Desvio Absoluto
        ws_kpi.cell(row=r, column=9).value = (
            f'=IF(OR(NOT(ISNUMBER(G{r})),NOT(ISNUMBER(H{r}))),"",H{r}-G{r})'
        )
        # J: Desvio %
        ws_kpi.cell(row=r, column=10).value = (
            f'=IF(OR(NOT(ISNUMBER(G{r})),NOT(ISNUMBER(H{r})),G{r}=0),"",I{r}/G{r})'
        )
        # K: Atingimento %
        ws_kpi.cell(row=r, column=11).value = (
            f'=IF(OR(NOT(ISNUMBER(G{r})),NOT(ISNUMBER(H{r})),G{r}=0),"",H{r}/G{r})'
        )
        # L: Farol
        ws_kpi.cell(row=r, column=12).value = (
            f'=IF(NOT(ISNUMBER(G{r})),"SEM META",'
            f'IF(G{r}=0,"SEM META",'
            f'IF(NOT(ISNUMBER(H{r})),"SEM APONTAMENTO",'
            f'IF(H{r}>=G{r},"VERDE",'
            f'IF(H{r}>=G{r}*0.9,"AMARELO","VERMELHO")))))'
        )

NEW_MAX = NEW_START + len(dates) * 3 - 1
print(f"3/5 CONTROLE KPI: +{len(dates)*3} linhas ({NEW_START}-{NEW_MAX}) GLOBO AÇOS")

# ─────────────────────────────────────────────────────────
# 4. PAINEL KPI — add GLOBO AÇOS + update ranges
# ─────────────────────────────────────────────────────────
ws_painel = wb['PAINEL KPI']

# Unmerge cells in the rows we need to modify (16+)
merges_to_remove = []
for mr in ws_painel.merged_cells.ranges:
    if mr.min_row >= 16:
        merges_to_remove.append(str(mr))
for mr_str in merges_to_remove:
    ws_painel.unmerge_cells(mr_str)

# Clear rows 16-18 (TOTAL was at 16, empty at 17, text at 18)
for r in range(16, 19):
    for c in range(1, 10):
        ws_painel.cell(row=r, column=c).value = None

# Write GLOBO AÇOS rows 16-18
ga_panel = [
    ('GLOBO AÇOS', 'PERFILAÇÃO'),
    ('GLOBO AÇOS', 'TELHAR'),
    ('GLOBO AÇOS', 'FECHAMENTO'),
]

for i, (empresa, servico) in enumerate(ga_panel):
    r = 16 + i
    for col in range(1, 8):
        ref = ws_painel.cell(row=15, column=col)
        tgt = ws_painel.cell(row=r, column=col)
        if ref.has_style:
            tgt.font = copy(ref.font)
            tgt.fill = copy(ref.fill)
            tgt.border = copy(ref.border)
            tgt.alignment = copy(ref.alignment)
            tgt.number_format = ref.number_format

    ws_painel.cell(row=r, column=1).value = empresa
    ws_painel.cell(row=r, column=2).value = servico
    ws_painel.cell(row=r, column=3).value = (
        f"=SUMIFS('CONTROLE KPI'!$G$2:$G${NEW_MAX},"
        f"'CONTROLE KPI'!$A$2:$A${NEW_MAX},\">=\"&$B$4,"
        f"'CONTROLE KPI'!$A$2:$A${NEW_MAX},\"<=\"&$B$5,"
        f"'CONTROLE KPI'!$C$2:$C${NEW_MAX},$A{r},"
        f"'CONTROLE KPI'!$D$2:$D${NEW_MAX},$B{r})"
    )
    ws_painel.cell(row=r, column=4).value = (
        f"=SUMIFS('CONTROLE KPI'!$H$2:$H${NEW_MAX},"
        f"'CONTROLE KPI'!$A$2:$A${NEW_MAX},\">=\"&$B$4,"
        f"'CONTROLE KPI'!$A$2:$A${NEW_MAX},\"<=\"&$B$5,"
        f"'CONTROLE KPI'!$C$2:$C${NEW_MAX},$A{r},"
        f"'CONTROLE KPI'!$D$2:$D${NEW_MAX},$B{r})"
    )
    ws_painel.cell(row=r, column=5).value = f'=D{r}-C{r}'
    ws_painel.cell(row=r, column=6).value = f'=IFERROR(D{r}/C{r},"")'
    ws_painel.cell(row=r, column=7).value = (
        f'=IF(C{r}="","SEM META",IF(C{r}=0,IF(D{r}>=0,"VERDE","VERMELHO"),'
        f'IF(D{r}>=C{r},"VERDE",IF(D{r}>=C{r}*0.9,"AMARELO","VERMELHO"))))'
    )

# TOTAL at row 19
ws_painel.merge_cells('A19:B19')
ws_painel.cell(row=19, column=1).value = 'TOTAL'
ws_painel.cell(row=19, column=3).value = '=SUM(C10:C18)'
ws_painel.cell(row=19, column=4).value = '=SUM(D10:D18)'
ws_painel.cell(row=19, column=5).value = '=D19-C19'
ws_painel.cell(row=19, column=6).value = '=IFERROR(D19/C19,"")'
ws_painel.cell(row=19, column=7).value = (
    '=IF(C19=0,"SEM META",IF(D19>=C19,"VERDE",'
    'IF(D19>=C19*0.9,"AMARELO","VERMELHO")))'
)

# Action text at row 21
ws_painel.merge_cells('A21:G21')
ws_painel.cell(row=21, column=1).value = (
    'Ação Lean: priorize linhas VERMELHAS; registre '
    'Categoria + Causa Raiz + Ação Corretiva no CONTROLE KPI.'
)

# Update ALL existing PAINEL KPI ranges from 2143 to NEW_MAX
ws_painel['E3'] = f"=SUMIFS('CONTROLE KPI'!$G$2:$G${NEW_MAX},'CONTROLE KPI'!$A$2:$A${NEW_MAX},$B$3)"
ws_painel['G3'] = f"=SUMIFS('CONTROLE KPI'!$H$2:$H${NEW_MAX},'CONTROLE KPI'!$A$2:$A${NEW_MAX},$B$3)"

ws_painel['G5'] = f"=COUNTIFS('CONTROLE KPI'!$A$2:$A${NEW_MAX},$B$3,'CONTROLE KPI'!$L$2:$L${NEW_MAX},\"VERMELHO\")"
ws_painel['I5'] = (
    f"=COUNTIFS('CONTROLE KPI'!$A$2:$A${NEW_MAX},\">=\"&$B$4,"
    f"'CONTROLE KPI'!$A$2:$A${NEW_MAX},\"<=\"&$B$5,"
    f"'CONTROLE KPI'!$I$2:$I${NEW_MAX},\"<0\","
    f"'CONTROLE KPI'!$M$2:$M${NEW_MAX},\"=\")"
)

# EJ rows 10-12
for r in range(10, 13):
    ws_painel.cell(row=r, column=3).value = (
        f"=SUMIFS('CONTROLE KPI'!$G$2:$G${NEW_MAX},"
        f"'CONTROLE KPI'!$A$2:$A${NEW_MAX},\">=\"&$B$4,"
        f"'CONTROLE KPI'!$A$2:$A${NEW_MAX},\"<=\"&$B$5,"
        f"'CONTROLE KPI'!$C$2:$C${NEW_MAX},$A{r},"
        f"'CONTROLE KPI'!$D$2:$D${NEW_MAX},$B{r})"
    )
    ws_painel.cell(row=r, column=4).value = (
        f"=SUMIFS('CONTROLE KPI'!$H$2:$H${NEW_MAX},"
        f"'CONTROLE KPI'!$A$2:$A${NEW_MAX},\">=\"&$B$4,"
        f"'CONTROLE KPI'!$A$2:$A${NEW_MAX},\"<=\"&$B$5,"
        f"'CONTROLE KPI'!$C$2:$C${NEW_MAX},$A{r},"
        f"'CONTROLE KPI'!$D$2:$D${NEW_MAX},$B{r})"
    )

# CMM rows 13-15
for r in range(13, 16):
    ws_painel.cell(row=r, column=3).value = (
        f"=SUMIFS('CONTROLE KPI'!$G$2:$G${NEW_MAX},"
        f"'CONTROLE KPI'!$A$2:$A${NEW_MAX},\">=\"&$B$4,"
        f"'CONTROLE KPI'!$A$2:$A${NEW_MAX},\"<=\"&$B$5,"
        f"'CONTROLE KPI'!$C$2:$C${NEW_MAX},$A{r},"
        f"'CONTROLE KPI'!$D$2:$D${NEW_MAX},$B{r})"
    )
    ws_painel.cell(row=r, column=4).value = (
        f"=SUMIFS('CONTROLE KPI'!$H$2:$H${NEW_MAX},"
        f"'CONTROLE KPI'!$A$2:$A${NEW_MAX},\">=\"&$B$4,"
        f"'CONTROLE KPI'!$A$2:$A${NEW_MAX},\"<=\"&$B$5,"
        f"'CONTROLE KPI'!$C$2:$C${NEW_MAX},$A{r},"
        f"'CONTROLE KPI'!$D$2:$D${NEW_MAX},$B{r})"
    )

print(f"4/5 PAINEL KPI: +3 GLOBO AÇOS (16-18), TOTAL→19, ranges até {NEW_MAX}")

# ─────────────────────────────────────────────────────────
# 5. Save
# ─────────────────────────────────────────────────────────
wb.save(DST)
print(f"\n5/5 Salvo: {DST}")
print(f"\nRESUMO FINAL:")
print(f"  BASE METAS DIÁRIAS: +9 colunas (U-AC) GLOBO AÇOS para {ws_bmd.max_row - 1} datas")
print(f"  PARAMETROS KPI: +3 linhas GLOBO AÇOS (11-13), ranges $5:$13")
print(f"  CONTROLE KPI: +{len(dates)*3} linhas ({NEW_START}-{NEW_MAX})")
print(f"  PAINEL KPI: +3 linhas GLOBO AÇOS (16-18), TOTAL→19, ranges→{NEW_MAX}")
